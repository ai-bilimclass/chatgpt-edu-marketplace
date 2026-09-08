#!/usr/bin/env python3
import argparse
import json
import re
import zipfile
from collections import Counter
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
SIZE = Pt(12)
CRITICAL_UNICODE = frozenset("ЁёҰұ")
EXPECTED_WEEKS = [8, 8, 10, 8]
EXPECTED_QUARTER_DATES = [
    ("2026-09-01", "2026-10-25"),
    ("2026-11-02", "2026-12-27"),
    ("2027-01-11", "2027-03-21"),
    ("2027-03-29", "2027-05-25"),
]
OBJECTIVE_CODE = re.compile(r"(?<!\d)(?:[5-9]|1[01])(?:\.\d+){3}(?!\d)")
OBJECTIVE_CODE_FORMAT_CHARS = "\u200b\u200c\u200d\u2060\ufeff"
COMPLEX_OBJECTIVE = re.compile(
    r"(?<!\w)(?:"
    r"талда\w*|зертте\w*|бағала\w*|дәлелде\w*|жобала\w*|құрастыр\w*|"
    r"эксперимент\w*|жоба\w*|өнім\w*|"
    r"анализ\w*|исслед\w*|оцен\w*|доказ\w*|обоснов\w*|проект\w*|"
    r"разработ\w*|созда\w*|эксперимент\w*|"
    r"analy[sz]\w*|investigat\w*|evaluat\w*|justif\w*|design\w*|"
    r"develop\w*|creat\w*|experiment\w*|project\w*"
    r")(?!\w)",
    re.IGNORECASE,
)
PLACEHOLDER = re.compile(
    r"(?:\bTBD\b|\bN/?A\b|уточнить|добавить\s+позже|заполнить\s+позже)",
    re.IGNORECASE,
)

SUBJECT_PROFILE_PATTERNS = {
    "languages_literature": re.compile(
        r"тілі|әдебиет|язык|литератур|language|literature|english|german|french|"
        r"ағылшын|неміс|француз|ұйғыр|өзбек|тәжік",
        re.IGNORECASE,
    ),
    "mathematics_informatics": re.compile(
        r"математ|алгебр|геометр|информат|mathemat|algebra|geometry|informatics|computer science",
        re.IGNORECASE,
    ),
    "natural_sciences": re.compile(
        r"жаратылыстану|естествозн|физик|хими|биологи|географ|natural science|"
        r"physics|chemistry|biology|geography", re.IGNORECASE
    ),
    "history_law": re.compile(
        r"тарих|истори|құқық|право|history|law", re.IGNORECASE
    ),
    "arts_technology": re.compile(
        r"музык|көркем еңбек|художественн|графика|жобалау|черчение|технолог|"
        r"music|art|graphics|design|drawing|technology", re.IGNORECASE
    ),
    "physical_aetd": re.compile(
        r"дене шынықтыру|физическ|алғашқы әскери|начальная военная|physical education|"
        r"initial military|basic military|аәtd|аәтд|нвтп",
        re.IGNORECASE,
    ),
    "applied_courses": re.compile(
        r"жаһандық құзыр|глобальн.*компет|кәсіпкерлік|предприниматель|бизнес|"
        r"global competenc|entrepreneurship|business",
        re.IGNORECASE,
    ),
}

SUMMATIVE_EXEMPT_SUBJECTS = (
    ({"музыка", "music"}, range(5, 7)),
    ({"художественный труд", "көркем еңбек", "artistic labor", "arts and crafts"}, range(5, 10)),
    ({"физическая культура", "дене шынықтыру", "physical education"}, range(5, 12)),
    (
        {
            "светскость и основы религиоведения",
            "зайырлылық және дінтану негіздері",
            "secularism and fundamentals of religious studies",
        },
        range(9, 10),
    ),
    (
        {
            "основы предпринимательства и бизнеса",
            "кәсіпкерлік және бизнес негіздері",
            "fundamentals of entrepreneurship and business",
        },
        range(10, 12),
    ),
    ({"графика и проектирование", "графика және жобалау", "graphics and design"}, range(10, 12)),
    (
        {
            "начальная военная подготовка",
            "начальная военная и технологическая подготовка",
            "нвп",
            "нвтп",
            "алғашқы әскери дайындық",
            "алғашқы әскери және технологиялық дайындық",
            "initial military training",
        },
        range(10, 12),
    ),
)


def normalize_policy_value(value: object) -> str:
    text = str(value or "").casefold().replace("ё", "е").replace("_", " ")
    text = re.sub(r"\([^)]*\)", " ", text)
    return " ".join(re.sub(r"[^\w]+", " ", text, flags=re.UNICODE).split())


def summative_assessment_exemption_reason(data: dict[str, Any]) -> str | None:
    academic_year = normalize_policy_value(data.get("academic_year", "2026–2027"))
    if academic_year != "2026 2027":
        return None
    grade = data.get("grade")
    if not isinstance(grade, int) or isinstance(grade, bool):
        return None
    component = normalize_policy_value(data.get("curriculum_component"))
    if component in {"variative", "вариативный", "вариативный компонент", "вариативтік компонент"}:
        if 5 <= grade <= 11:
            return "вариативный компонент, 5–11 классы"
    subject = normalize_policy_value(data.get("subject"))
    for aliases, grades in SUMMATIVE_EXEMPT_SUBJECTS:
        if subject in aliases and grade in grades:
            return f"{data.get('subject')}, {grade} класс"
    return None


def apply_summative_assessment_exemption(data: dict[str, Any], reason: str) -> None:
    marker = re.compile(r"(?<!\w)(?:БЖБ|ТЖБ|СОР|СОЧ|SAU|SAT)(?:\s*№?\s*\d+)?(?!\w)", re.IGNORECASE)
    for quarter in data.get("quarters", []):
        quarter["soch_required"] = False
        for lesson in quarter.get("lessons", []):
            topic = str(lesson.get("topic", ""))
            if marker.search(topic):
                raise ValueError(
                    "The teacher source contains a separate summative-assessment lesson for "
                    f"an exempt subject ({reason}); resolve the source conflict before generation"
                )
            lesson["assessment_type"] = ""
    data["summative_assessment_exempt"] = True
    data["summative_assessment_exemption_reason"] = reason
PROFILE_DIRECTION_SUBJECT = re.compile(
    r"алгебр|геометр|информат|физик|хими|биологи|географ|дүниежүзі тарих|"
    r"всемирн.*истори|құқық|право|algebra|geometry|informatics|computer science|"
    r"physics|chemistry|biology|geography|world history|law",
    re.IGNORECASE,
)
PRACTICAL_MODES = {"laboratory", "experiment", "practical", "field_training"}


def infer_subject_profile(subject: str) -> str:
    matches = [
        profile for profile, pattern in SUBJECT_PROFILE_PATTERNS.items()
        if pattern.search(subject)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Пәндік профиль анықталмады: «{subject}». Пәннің ресми атауын нақтылаңыз."
        )
    return matches[0]


def validate_subject_context(data: dict[str, Any]) -> None:
    grade = int(data["grade"])
    expected_level = "basic_secondary" if grade <= 9 else "general_secondary"
    if data["education_level"] != expected_level:
        raise ValueError(f"education_level must be {expected_level} for grade {grade}")
    if str(data["instruction_language"]).strip() not in {"ru", "kk", "en", "mixed"}:
        raise ValueError("instruction_language must be ru, kk, en, or mixed")
    if "399" not in str(data["curriculum_order"]):
        raise ValueError("curriculum_order must identify the applicable order №399")
    if not str(data["curriculum_appendix"]).strip():
        raise ValueError("curriculum_appendix must identify the verified №399 appendix")
    try:
        date.fromisoformat(str(data["curriculum_revision_date"]))
    except ValueError as exc:
        raise ValueError("curriculum_revision_date must use YYYY-MM-DD") from exc
    expected_profile = infer_subject_profile(str(data["subject"]))
    if data["subject_profile"] != expected_profile:
        raise ValueError(
            f"subject_profile must be {expected_profile} for {data['subject']}"
        )
    direction = str(data["profile_direction"]).strip()
    if grade >= 10 and PROFILE_DIRECTION_SUBJECT.search(str(data["subject"])):
        if direction not in {"ЖМБ", "ҚГБ"}:
            raise ValueError("10–11 сыныптың профильдік пәні үшін ЖМБ немесе ҚГБ таңдаңыз")
    elif direction != "not_applicable":
        raise ValueError("profile_direction must be not_applicable for this subject and grade")

    for quarter in data.get("quarters", []):
        for lesson in quarter.get("lessons", []):
            mode = str(lesson.get("lesson_mode", "theory")).strip()
            if expected_profile in {"natural_sciences", "physical_aetd"} and mode in PRACTICAL_MODES:
                if not str(lesson.get("practical_work", "")).strip():
                    raise ValueError("Практикалық/зертханалық сабаққа practical_work атауы қажет")
                if lesson.get("safety_required") is not True:
                    raise ValueError("Практикалық/зертханалық сабаққа safety_required: true қажет")

TEXT = {
    "ru": {
        "title": "Среднесрочный (календарно-тематический) план по предмету",
        "class": "класс",
        "total": "Итого",
        "hours": "часов",
        "per_week": "в неделю",
        "year": "Учебный год",
        "school_week": "Учебная неделя",
        "calendar_note": "Сроки обучения и каникул утверждены Министерством просвещения РК на 2026–2027 учебный год; праздники и переносы проверяются отдельно.",
        "headers": ["№ п/п", "Раздел /\nСквозные темы", "Тема урока", "Цели обучения", "Количество часов", "Сроки", "Примечание"],
        "quarter": "четверть",
        "weeks": "недель",
    },
    "kk": {
        "title": "Пән бойынша орта мерзімді (күнтізбелік-тақырыптық) жоспар",
        "class": "сынып",
        "total": "Барлығы",
        "hours": "сағат",
        "per_week": "аптасына",
        "year": "Оқу жылы",
        "school_week": "Оқу аптасы",
        "calendar_note": "2026–2027 оқу жылындағы оқу және каникул мерзімдерін ҚР Оқу-ағарту министрлігі бекітті; мерекелер мен демалыс күндерінің ауыстырылуы бөлек тексеріледі.",
        "headers": ["№", "Бөлім /\nОртақ тақырыптар", "Сабақтың тақырыбы", "Оқу мақсаттары", "Сағат саны", "Мерзімі", "Ескерту"],
        "quarter": "тоқсан",
        "weeks": "апта",
    },
    "en": {
        "title": "Medium-term (calendar-thematic) subject plan",
        "class": "grade",
        "total": "Total",
        "hours": "hours",
        "per_week": "per week",
        "year": "Academic year",
        "school_week": "School week",
        "calendar_note": "The Ministry of Education of Kazakhstan approved the 2026–2027 study and holiday dates; public holidays and transferred days off are verified separately.",
        "headers": ["No.", "Section /\nCross-cutting themes", "Lesson topic", "Learning objectives", "Number of hours", "Dates", "Notes"],
        "quarter": "term",
        "weeks": "weeks",
    },
}


def trilingual(key: str) -> str:
    return " / ".join(TEXT[lang][key] for lang in ("kk", "ru", "en"))


def labels(language: str) -> dict[str, Any]:
    if language in TEXT:
        return TEXT[language]
    if language == "trilingual":
        return {
            "title": trilingual("title"),
            "class": trilingual("class"),
            "total": trilingual("total"),
            "hours": trilingual("hours"),
            "per_week": trilingual("per_week"),
            "year": trilingual("year"),
            "school_week": trilingual("school_week"),
            "calendar_note": "\n".join(TEXT[lang]["calendar_note"] for lang in ("kk", "ru", "en")),
            "headers": [" / ".join(TEXT[lang]["headers"][i] for lang in ("kk", "ru", "en")) for i in range(7)],
            "quarter": trilingual("quarter"),
            "weeks": trilingual("weeks"),
        }
    raise ValueError("language must be ru, kk, en, or trilingual")


def school_week_text(value: Any, language: str) -> str:
    normalized = str(value).strip().lower()
    translations = {
        "5-day": {
            "ru": "пятидневная",
            "kk": "бес күндік",
            "en": "five-day",
        },
        "6-day": {
            "ru": "шестидневная",
            "kk": "алты күндік",
            "en": "six-day",
        },
    }
    if normalized not in translations:
        return str(value)
    if language == "trilingual":
        return " / ".join(translations[normalized][lang] for lang in ("kk", "ru", "en"))
    return translations[normalized][language]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def apply_font_to_paragraph(paragraph, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.size = SIZE
        run.font.color.rgb = RGBColor(0, 0, 0)
        if bold is not None:
            run.bold = bold
        r_pr = run._element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        for channel in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{channel}"), FONT)


def validate_unicode_value(value: Any, path: str = "data") -> None:
    if isinstance(value, str):
        if "\ufffd" in value:
            raise ValueError(f"{path} contains the Unicode replacement character U+FFFD")
        try:
            value.encode("utf-8", "strict")
        except UnicodeEncodeError as exc:
            raise ValueError(f"{path} contains an unpaired Unicode surrogate") from exc
        invalid = [char for char in value if ord(char) < 32 and char not in "\t\n\r"]
        if invalid:
            raise ValueError(f"{path} contains a prohibited control character")
        return
    if isinstance(value, dict):
        for key, item in value.items():
            validate_unicode_value(item, f"{path}.{key}")
    elif isinstance(value, list):
        for index, item in enumerate(value):
            validate_unicode_value(item, f"{path}[{index}]")


def normalize_objective_code_formatting(value: str) -> str:
    """Remove spreadsheet formatting characters only inside numeric objective codes."""
    escaped = re.escape(OBJECTIVE_CODE_FORMAT_CHARS)
    pattern = re.compile(
        rf"(?<=\d)[{escaped}]+(?=[.\d])|(?<=\.)[{escaped}]+(?=\d)"
    )
    previous = None
    normalized = value
    while normalized != previous:
        previous = normalized
        normalized = pattern.sub("", normalized)
    return normalized


def normalize_objective_codes_in_data(data: dict[str, Any]) -> None:
    for item in data.get("extended_objectives", []):
        if isinstance(item, dict) and isinstance(item.get("learning_objective"), str):
            item["learning_objective"] = normalize_objective_code_formatting(
                item["learning_objective"]
            )
    for quarter in data.get("quarters", []):
        for lesson in quarter.get("lessons", []):
            if isinstance(lesson.get("learning_objectives"), str):
                lesson["learning_objectives"] = normalize_objective_code_formatting(
                    lesson["learning_objectives"]
                )


def expected_critical_unicode(data: dict[str, Any]) -> set[str]:
    displayed: list[str] = [str(data.get("subject", ""))]
    for quarter in data.get("quarters", []):
        displayed.append(str(quarter.get("name", "")))
        for lesson in quarter.get("lessons", []):
            displayed.extend(str(lesson.get(key, "")) for key in (
                "section", "topic", "learning_objectives", "assessment_type", "note"
            ))
    return {char for text in displayed for char in text if char in CRITICAL_UNICODE}


def set_cell_text(cell, text: Any, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = "" if text is None else str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.keep_together = True
        apply_font_to_paragraph(paragraph, bold=bold)


def assessment_terms(language: str) -> tuple[str, str]:
    if language == "kk":
        return "БЖБ", "ТЖБ"
    if language == "ru":
        return "СОР", "СОЧ"
    if language == "en":
        return "SAU", "SAT"
    raise ValueError(
        "Үштілді КТП-дағы жиынтық бағалау белгілері қай тілде көрсетілсін: "
        "қазақша БЖБ/ТЖБ, орысша СОР/СОЧ немесе ағылшынша SAU/SAT?"
    )


def apply_automatic_soch(quarter: dict[str, Any], language: str) -> None:
    """Normalize quarter assessment marking in the document language."""
    _, term_label = assessment_terms(language)
    soch_required = quarter.get("soch_required")
    if not isinstance(soch_required, bool):
        raise ValueError("Every quarter requires boolean soch_required")

    lessons = quarter.get("lessons", [])
    assessment_types = []
    for lesson in lessons:
        value = str(lesson.get("assessment_type", "")).strip().upper()
        # Quarter- and section-level flags are authoritative. Preserve only a
        # possible СОЧ marker here; БЖБ/СОР is rebuilt from section boundaries.
        normalized = (
            term_label
            if re.search(r"(?<!\w)(?:ТЖБ|СОЧ|SAT)(?!\w)", value)
            else ""
        )
        lesson["assessment_type"] = normalized
        assessment_types.append(normalized)
    if any(
        re.search(r"(?<!\w)(?:БЖБ|СОР|SAU)(?!\w)", str(lesson.get("topic", "")), re.IGNORECASE)
        for lesson in lessons
    ):
        raise ValueError("Separate БЖБ/СОР/SAU lesson rows are prohibited; marking is automatic")

    if not soch_required:
        if any(value == term_label for value in assessment_types):
            raise ValueError("A quarter with soch_required=false must not contain a ТЖБ/СОЧ row")
        return

    if len(lessons) < 2:
        raise ValueError("A quarter requiring СОЧ must contain at least two lessons")

    # The quarter-level flag is authoritative. Remove stale or duplicate ТЖБ/СОЧ
    # marks from the input and insert exactly one on the penultimate lesson.
    for lesson, assessment_type in zip(lessons, assessment_types):
        if assessment_type == term_label:
            lesson["assessment_type"] = ""
    soch_lesson = lessons[-2]
    soch_lesson["assessment_type"] = term_label
    soch_lesson["topic"] = ""
    soch_lesson["learning_objectives"] = ""


def apply_section_assessments(quarters: list[dict[str, Any]], language: str) -> None:
    """Place БЖБ/СОР at section ends, always before quarter-end ТЖБ/СОЧ."""
    section_label, _ = assessment_terms(language)
    flattened = [lesson for quarter in quarters for lesson in quarter.get("lessons", [])]
    global_positions = {id(lesson): index for index, lesson in enumerate(flattened)}
    for quarter in quarters:
        quarter_lessons = quarter.get("lessons", [])
        soch_index = len(quarter_lessons) - 2 if quarter.get("soch_required") else None
        number = 1
        for local_index, lesson in enumerate(quarter_lessons):
            position = global_positions[id(lesson)]
            next_lesson = flattened[position + 1] if position + 1 < len(flattened) else None
            current_section = str(lesson.get("section", "")).strip()
            next_section = str(next_lesson.get("section", "")).strip() if next_lesson else None
            if next_lesson is not None and next_section == current_section:
                continue
            target_index = local_index
            if soch_index is not None and target_index >= soch_index:
                target_index = next(
                    (
                        candidate
                        for candidate in range(soch_index - 1, -1, -1)
                        if str(quarter_lessons[candidate].get("section", "")).strip()
                        == current_section
                    ),
                    -1,
                )
                if target_index < 0:
                    quarter_name = str(quarter.get("name", "Тоқсан")).strip()
                    raise ValueError(
                        f"{quarter_name} ішінде {current_section} бөлімінің БЖБ/СОР-ын "
                        "ТЖБ/СОЧ-қа дейінгі қай сабаққа орналастырайық?"
                    )
            label = f"{section_label} №{number}"
            quarter_lessons[target_index]["assessment_type"] = label
            number += 1


def validate_assessment_order(quarters: list[dict[str, Any]], language: str) -> None:
    """Reject misplaced or duplicated quarter-end assessment markers."""
    section_label, term_label = assessment_terms(language)
    for quarter in quarters:
        if not quarter.get("soch_required"):
            continue
        lessons = quarter.get("lessons", [])
        soch_positions = [
            index
            for index, lesson in enumerate(lessons)
            if str(lesson.get("assessment_type", "")).strip() == term_label
        ]
        expected_soch = len(lessons) - 2
        if soch_positions != [expected_soch]:
            raise ValueError("ТЖБ/СОЧ must appear exactly once on the penultimate lesson")
        soch_lesson = lessons[expected_soch]
        if str(soch_lesson.get("topic", "")).strip() or str(
            soch_lesson.get("learning_objectives", "")
        ).strip():
            raise ValueError("ТЖБ/СОЧ lesson must not contain a topic or learning objective")
        if float(soch_lesson.get("hours", soch_lesson.get("hours_total", 0)) or 0) != 1:
            raise ValueError("ТЖБ/СОЧ must remain one existing lesson hour")
        if any(
            re.search(
                rf"(?<!\w){re.escape(section_label)}(?!\w)",
                str(lesson.get("assessment_type", "")),
                re.IGNORECASE,
            )
            for lesson in lessons[expected_soch:]
        ):
            raise ValueError("БЖБ/СОР must be placed before ТЖБ/СОЧ")
        if str(lessons[-1].get("assessment_type", "")).strip():
            raise ValueError("The final lesson of a quarter must not contain an assessment marker")


EXTERNAL_URL = re.compile(r"(?:https?://|www\.)", re.IGNORECASE)


def reject_external_urls(value: Any, path: str = "input") -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            if str(key).casefold() in {"url", "source_url", "external_url"}:
                raise ValueError(f"External URL fields are forbidden in KTP input: {path}.{key}")
            reject_external_urls(child, f"{path}.{key}")
    elif isinstance(value, list):
        for index, child in enumerate(value):
            reject_external_urls(child, f"{path}[{index}]")
    elif isinstance(value, str) and EXTERNAL_URL.search(value):
        raise ValueError(f"External URLs are forbidden in KTP input: {path}")


def validate(data: dict[str, Any]) -> None:
    reject_external_urls(data)
    normalize_objective_codes_in_data(data)
    validate_unicode_value(data)
    required = [
        "subject", "grade", "education_level", "instruction_language", "profile_direction",
        "curriculum_order", "curriculum_appendix", "curriculum_revision_date", "subject_profile",
        "language", "requested_language", "objectives_language",
        "language_mismatch_confirmed", "academic_year", "total_hours",
        "hours_per_week", "lesson_weekdays",
        "content_source", "content_source_type", "teacher_program_file",
        "official_program_verified", "source_content_complete",
        "source_conflicts",
        "calendar_verified", "calendar_source", "calendar_source_type", "calendar_source_complete",
        "non_instruction_dates", "public_holidays", "quarters"
    ]
    missing = [k for k in required if k not in data]
    if missing:
        raise ValueError(f"Missing required fields: {', '.join(missing)}")
    if int(data["grade"]) not in range(5, 12):
        raise ValueError("Grade must be from 5 to 11")
    validate_subject_context(data)
    if str(data["academic_year"]).replace("-", "–") != "2026–2027":
        raise ValueError("Academic year must be 2026–2027")
    labels(str(data["language"]))
    requested_language = str(data["requested_language"]).strip()
    objectives_language = str(data["objectives_language"]).strip()
    if requested_language not in {"ru", "kk", "en", "trilingual"}:
        raise ValueError("requested_language must be ru, kk, en, or trilingual")
    if objectives_language not in {"ru", "kk", "en", "mixed"}:
        raise ValueError("objectives_language must be ru, kk, en, or mixed")
    if not isinstance(data["language_mismatch_confirmed"], bool):
        raise ValueError("language_mismatch_confirmed must be boolean")
    instruction_language = str(data["instruction_language"]).strip()
    if requested_language != instruction_language:
        raise ValueError("requested_language must mirror instruction_language; do not ask document language separately")
    output_language = str(data["language"]).strip()
    explicit_override = data["language_mismatch_confirmed"] is True
    if output_language != instruction_language and not explicit_override:
        raise ValueError("Output language must match instruction_language unless the teacher explicitly requested another document language")
    if "color_theme" in data and data["color_theme"] != {"mode": "official-monochrome"}:
        raise ValueError("Only official monochrome formatting is allowed")
    if data["official_program_verified"] is not True:
        raise ValueError("The official program for the selected subject and grade must be verified")
    if data["content_source_type"] != "teacher_uploaded_program":
        raise ValueError("KTP generation requires a teacher-uploaded curriculum program")
    if not str(data["teacher_program_file"]).strip():
        raise ValueError("teacher_program_file must identify the curriculum file uploaded by the teacher")
    if data["source_content_complete"] is not True:
        raise ValueError("All official sections, works/topics, and objectives must be extracted before generation")
    for field in ("subject", "content_source", "teacher_program_file"):
        if not str(data.get(field, "")).strip():
            raise ValueError(f"{field} must not be empty")
    source_conflicts = data["source_conflicts"]
    if not isinstance(source_conflicts, list):
        raise ValueError("source_conflicts must be a list")
    for item in source_conflicts:
        if not isinstance(item, dict) or any(
            not str(item.get(key, "")).strip()
            for key in ("location", "type", "source_text", "resolution")
        ):
            raise ValueError(
                "Every source conflict requires location, type, source_text, and resolution"
            )
    extended_objectives = data.get("extended_objectives", [])
    if not isinstance(extended_objectives, list):
        raise ValueError("extended_objectives must be a list when provided")
    for item in extended_objectives:
        if not isinstance(item, dict) or not str(item.get("learning_objective", "")).strip():
            raise ValueError("Every extended_objectives item requires learning_objective")
        try:
            lesson_count = int(item["lesson_count"])
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError("Every extended_objectives item requires integer lesson_count") from exc
        if lesson_count <= 1:
            raise ValueError("extended_objectives lesson_count must be greater than 1")
    try:
        hours_per_week = int(data["hours_per_week"])
    except (TypeError, ValueError) as exc:
        raise ValueError("hours_per_week must be an integer greater than or equal to 1") from exc
    if float(data["hours_per_week"]) != hours_per_week or hours_per_week < 1:
        raise ValueError("hours_per_week must be an integer greater than or equal to 1")
    quarters = data["quarters"]
    weekdays = data["lesson_weekdays"]
    if not isinstance(weekdays, list) or not weekdays:
        raise ValueError("lesson_weekdays must contain at least one scheduled lesson day")
    weekdays = [int(value) for value in weekdays]
    if len(weekdays) != hours_per_week:
        raise ValueError(
            "lesson_weekdays must contain exactly one weekly slot per hour_per_week; repeated days are allowed"
        )
    if any(value not in range(7) for value in weekdays):
        raise ValueError("lesson_weekdays values must be integers from 0 (Monday) to 6 (Sunday)")
    school_week = str(data.get("school_week", "5-day")).strip().lower()
    if school_week not in {"5-day", "6-day"}:
        raise ValueError("school_week must be 5-day or 6-day when provided")
    if school_week == "5-day" and any(value > 4 for value in weekdays):
        raise ValueError("A five-day school week cannot schedule lessons on Saturday or Sunday")
    if school_week == "6-day" and any(value > 5 for value in weekdays):
        raise ValueError("A six-day school week cannot schedule lessons on Sunday")
    if data.get("calendar_verified") is not True:
        raise ValueError("Calendar must be verified before mandatory lesson dates can be generated")
    if not str(data.get("calendar_source", "")).strip():
        raise ValueError("calendar_source is required")
    if data["calendar_source_type"] not in {"builtin_approved_calendar_2026_2027", "teacher_uploaded_calendar"}:
        raise ValueError("calendar_source_type must use the built-in approved calendar or a teacher-uploaded calendar")
    if data["calendar_source_complete"] is not True:
        raise ValueError("Missing mandatory calendar data; request a source from the teacher and stop generation")
    try:
        excluded_dates = {date.fromisoformat(str(value)) for value in data["non_instruction_dates"]}
    except (TypeError, ValueError) as exc:
        raise ValueError("non_instruction_dates must be a list of YYYY-MM-DD dates") from exc
    if not isinstance(data["public_holidays"], list):
        raise ValueError("public_holidays must be a list")
    holiday_dates: set[date] = set()
    for item in data["public_holidays"]:
        if not isinstance(item, dict) or not str(item.get("name", "")).strip():
            raise ValueError("Every public holiday requires date and name")
        try:
            holiday_date = date.fromisoformat(str(item["date"]))
            transfer_value = item.get("official_transfer_date")
            if transfer_value not in (None, ""):
                transfer_date = date.fromisoformat(str(transfer_value))
                if transfer_date in excluded_dates:
                    raise ValueError("official_transfer_date cannot be a non-instruction date")
            teacher_value = item.get("teacher_transfer_date")
            if teacher_value not in (None, ""):
                teacher_date = date.fromisoformat(str(teacher_value))
                if teacher_date in excluded_dates:
                    raise ValueError("teacher_transfer_date cannot be a non-instruction date")
        except (KeyError, ValueError) as exc:
            raise ValueError("Holiday dates must use YYYY-MM-DD") from exc
        if holiday_date not in excluded_dates:
            raise ValueError("Every public holiday date must also be in non_instruction_dates")
        if holiday_date in holiday_dates:
            raise ValueError("public_holidays must not contain duplicate dates")
        holiday_dates.add(holiday_date)
    if len(quarters) != 4:
        raise ValueError("Exactly four quarters are required")
    actual_weeks = [int(q.get("weeks", 0)) for q in quarters]
    if actual_weeks != EXPECTED_WEEKS:
        raise ValueError(f"Quarter durations must be {EXPECTED_WEEKS}, got {actual_weeks}")
    for index, quarter in enumerate(quarters):
        try:
            start = date.fromisoformat(str(quarter["start_date"]))
            end = date.fromisoformat(str(quarter["end_date"]))
        except (KeyError, ValueError) as exc:
            raise ValueError("Every quarter requires valid start_date and end_date in YYYY-MM-DD format") from exc
        if start > end:
            raise ValueError("Quarter start_date cannot be after end_date")
        expected_start, expected_end = EXPECTED_QUARTER_DATES[index]
        if (start.isoformat(), end.isoformat()) != (expected_start, expected_end):
            raise ValueError(
                f"Quarter {index + 1} dates must be {expected_start}–{expected_end}, "
                f"got {start.isoformat()}–{end.isoformat()}"
            )
    lessons = [lesson for q in quarters for lesson in q.get("lessons", [])]
    if not lessons:
        raise ValueError("At least one lesson is required")
    expected_numbers = list(range(1, len(lessons) + 1))
    actual_numbers = [int(lesson.get("number", 0)) for lesson in lessons]
    if actual_numbers != expected_numbers:
        raise ValueError(
            f"Lesson numbering must be sequential from 1 to {len(lessons)}, got {actual_numbers}"
        )
    for item in extended_objectives:
        requested = str(item["learning_objective"]).strip()
        codes = OBJECTIVE_CODE.findall(requested)
        if not codes:
            raise ValueError("Every extended objective must contain an exact learning-objective code")
        matching_rows = sum(
            1 for lesson in lessons
            if all(code in str(lesson.get("learning_objectives", "")) for code in codes)
        )
        if matching_rows < int(item["lesson_count"]):
            raise ValueError(
                f"Extended objective {', '.join(codes)} requires {item['lesson_count']} separate lesson rows, "
                f"but only {matching_rows} were found"
            )
    exemption_reason = summative_assessment_exemption_reason(data)
    if exemption_reason:
        apply_summative_assessment_exemption(data, exemption_reason)
    for quarter in quarters:
        quarter_lessons = quarter.get("lessons", [])
        quarter_hours = sum(
            float(x.get("hours", x.get("hours_total", 0)) or 0)
            for x in quarter_lessons
        )
        expected_quarter_hours = int(quarter["weeks"]) * hours_per_week
        if abs(quarter_hours - expected_quarter_hours) > 1e-9:
            raise ValueError(
                f"{quarter.get('name', 'Quarter')} hours ({quarter_hours:g}) "
                f"must equal weeks × hours_per_week ({expected_quarter_hours:g})"
            )
        if len(quarter_lessons) != int(expected_quarter_hours):
            raise ValueError(
                f"{quarter.get('name', 'Quarter')} must contain {int(expected_quarter_hours)} "
                "separate one-hour lesson rows"
            )
        if not exemption_reason:
            apply_automatic_soch(quarter, str(data["language"]))
    if not exemption_reason:
        apply_section_assessments(quarters, str(data["language"]))
        validate_assessment_order(quarters, str(data["language"]))
    _, term_label = assessment_terms(str(data["language"]))
    for lesson in lessons:
        if not str(lesson.get("section", "")).strip():
            raise ValueError("Every lesson must have a section or cross-cutting theme")
        hours = float(lesson.get("hours", lesson.get("hours_total", 0)) or 0)
        if hours != 1:
            raise ValueError("Every generated lesson row must represent exactly one hour")
        if str(lesson.get("assessment_type", "")).strip() == term_label:
            continue
        topic = str(lesson.get("topic", "")).strip()
        if not topic:
            raise ValueError(
                "KTP generation is prohibited: every lesson must have a methodically distributed topic"
            )
        objectives = str(lesson.get("learning_objectives", "")).strip()
        if not objectives:
            raise ValueError(
                "KTP generation is prohibited: every lesson requires an exact official learning objective"
            )
        if PLACEHOLDER.search(topic) or PLACEHOLDER.search(objectives):
            raise ValueError(
                "KTP generation is prohibited: topics and learning objectives may not contain placeholders"
            )
        if not OBJECTIVE_CODE.search(objectives):
            raise ValueError(
                "KTP generation is prohibited: every learning objective must contain an exact code"
            )
        objective_codes = OBJECTIVE_CODE.findall(objectives)
        if len(objective_codes) > 2:
            raise ValueError(
                "Оқу мақсаттары бір сабаққа сыймайды: бір сабаққа ең көбі екі мақсат. "
                "Қай мақсатты басқа сабаққа ауыстырайық?"
            )
        high_load = (
            COMPLEX_OBJECTIVE.search(objectives) is not None
            or len(objectives) > 360
            or len(topic) > 180
        )
        if len(objective_codes) == 2 and high_load:
            raise ValueError(
                "Күрделі немесе көлемді оқу мақсаты басқа мақсатпен бір сабаққа "
                "біріктірілмейді. Қай мақсатты басқа сабаққа ауыстырайық?"
            )
        wording = OBJECTIVE_CODE.sub("", objectives)
        wording = re.sub(r"[\s,;:–—-]+", " ", wording).strip()
        if len(wording) < 8:
            raise ValueError(
                "KTP generation is prohibited: every learning-objective code must include its full wording"
            )
    actual = sum(float(x.get("hours", x.get("hours_total", 0)) or 0) for x in lessons)
    expected = float(data["total_hours"])
    if abs(actual - expected) > 1e-9:
        raise ValueError(f"Lesson hours ({actual:g}) do not equal total hours ({expected:g})")
    expected_year = sum(EXPECTED_WEEKS) * hours_per_week
    if abs(expected - expected_year) > 1e-9:
        raise ValueError(
            f"Total hours ({expected:g}) must equal 34 weeks × hours_per_week ({expected_year:g})"
        )


def holiday_note(language: str, name: str, original: date, target: date, method: str) -> str:
    method_text = {
        "official": {"ru": "официальный перенос", "kk": "ресми ауыстыру", "en": "official transfer"},
        "next": {"ru": "объединено со следующим уроком", "kk": "келесі сабақпен біріктірілді", "en": "combined with the next lesson"},
        "previous": {"ru": "объединено с предыдущим уроком", "kk": "алдыңғы сабақпен біріктірілді", "en": "combined with the previous lesson"},
        "teacher": {"ru": "дата подтверждена учителем", "kk": "күнді мұғалім растады", "en": "date confirmed by the teacher"},
    }
    templates = {
        "ru": "Праздничный день {original}: {name}; перенесено на {target}, {method}",
        "kk": "Мереке күні {original}: {name}; {target} күніне ауыстырылды, {method}",
        "en": "Public holiday {original}: {name}; moved to {target}, {method}",
    }
    languages = ("kk", "ru", "en") if language == "trilingual" else (language,)
    return " / ".join(templates[lang].format(
        original=original.strftime("%d.%m.%Y"),
        name=name,
        target=target.strftime("%d.%m.%Y"),
        method=method_text[method][lang],
    ) for lang in languages)


def transfer_date_question(language: str, name: str, original: date, quarter_name: str) -> str:
    templates = {
        "ru": "На какую дату в {quarter} перенести урок от {original} в связи с праздником «{name}»?",
        "kk": "{name} мерекесіне байланысты {original} күнгі сабақты {quarter} ішінде қай күнге ауыстырайық?",
        "en": "To which date in {quarter} should the {original} lesson be moved because of the {name} holiday?",
    }
    languages = ("kk", "ru", "en") if language == "trilingual" else (language,)
    return " / ".join(templates[lang].format(
        name=name, original=original.strftime("%d.%m.%Y"), quarter=quarter_name
    ) for lang in languages)


def assign_lesson_dates(data: dict[str, Any]) -> None:
    weekday_slots = Counter(int(value) for value in data["lesson_weekdays"])
    excluded = {date.fromisoformat(str(value)) for value in data["non_instruction_dates"]}
    holidays = {date.fromisoformat(str(item["date"])): item for item in data["public_holidays"]}
    for quarter in data["quarters"]:
        start = date.fromisoformat(str(quarter["start_date"]))
        end = date.fromisoformat(str(quarter["end_date"]))
        current = start
        planned_slots: list[date] = []
        while current <= end:
            planned_slots.extend([current] * weekday_slots.get(current.weekday(), 0))
            current += timedelta(days=1)
        valid_slots = [value for value in planned_slots if value not in excluded]
        assignments: list[tuple[date, str]] = []
        for planned in planned_slots:
            if planned not in excluded:
                assignments.append((planned, ""))
                continue
            holiday = holidays.get(planned)
            if holiday is None:
                continue
            transfer_value = holiday.get("official_transfer_date")
            target = date.fromisoformat(str(transfer_value)) if transfer_value not in (None, "") else None
            method = "official"
            teacher_value = holiday.get("teacher_transfer_date")
            if target is None and teacher_value not in (None, ""):
                target = date.fromisoformat(str(teacher_value))
                method = "teacher"
            if target is None or not (start <= target <= end):
                later = next((value for value in valid_slots if value > planned), None)
                earlier = next((value for value in reversed(valid_slots) if value < planned), None)
                target = later or earlier
                method = "next" if later is not None else "previous"
            if target is None:
                raise ValueError(transfer_date_question(
                    str(data["language"]), str(holiday["name"]), planned,
                    str(quarter.get("name", "Quarter"))
                ))
            assignments.append((target, holiday_note(
                str(data["language"]), str(holiday["name"]), planned, target, method
            )))
        needed = sum(int(float(x.get("hours", x.get("hours_total", 0)))) for x in quarter["lessons"])
        if len(assignments) < needed:
            # A quarter can contain the official number of teaching weeks yet
            # begin after one of this subject's weekly slots. Preserve the
            # official hour total by combining the missing opening slot with
            # the first actual lesson date and make that visible in the KTP.
            if not assignments:
                raise ValueError("No valid lesson dates are available in the quarter")
            missing = needed - len(assignments)
            first_date = assignments[0][0]
            note = (
                "Оқу тоқсаны аптаның ортасында басталуына байланысты "
                "біріктірілген сабақ"
            )
            assignments = [(first_date, note)] * missing + assignments
        position = 0
        for lesson in quarter["lessons"]:
            count = int(float(lesson.get("hours", lesson.get("hours_total", 0))))
            selected = assignments[position:position + count]
            lesson["date"] = ", ".join(value.strftime("%d.%m.%Y") for value, _ in selected)
            transfer_notes = [note for _, note in selected if note]
            if transfer_notes:
                existing_note = str(lesson.get("note", "")).strip()
                lesson["note"] = " — ".join(value for value in (existing_note, *transfer_notes) if value)
            position += count


def add_centered(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text).bold = bold
    apply_font_to_paragraph(p)



def safe_filename_part(value: Any) -> str:
    text = re.sub(r'[\\/:*?"<>|]+', ' ', str(value)).strip()
    text = re.sub(r'\s+', ' ', text)
    return text or "Предмет"


def default_output_name(data: dict[str, Any]) -> str:
    grade = int(data["grade"])
    subject = safe_filename_part(data["subject"])
    return f"{grade} класс КТП {subject}.docx"


def add_plan_table(doc: Document, lab: dict[str, Any]):
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    # Keep the lesson-number column wide enough for large dynamic annual loads.
    widths = [Cm(1.5), Cm(3.4), Cm(4.4), Cm(7.2), Cm(2.0), Cm(2.4), Cm(3.0)]
    for i, width in enumerate(widths):
        table.columns[i].width = width
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = width

    header = table.rows[0]
    for i, text in enumerate(lab["headers"]):
        set_cell_text(header.cells[i], text, True, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(header.cells[i], "FFFFFF")
    set_repeat_table_header(header)
    header.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    prevent_row_split(header)
    return table


def build(data: dict[str, Any], output: Path) -> None:
    validate(data)
    assign_lesson_dates(data)
    lab = labels(str(data["language"]))
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    style_r_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for channel in ("ascii", "hAnsi", "eastAsia", "cs"):
        style_r_fonts.set(qn(f"w:{channel}"), FONT)
    language_codes = {"ru": "ru-RU", "kk": "kk-KZ", "en": "en-US", "trilingual": "kk-KZ"}
    style_language = OxmlElement("w:lang")
    style_language.set(qn("w:val"), language_codes[str(data["language"])])
    style._element.get_or_add_rPr().append(style_language)

    add_centered(doc, lab["title"], True)
    add_centered(doc, f"{data['subject']} — {data['grade']} {lab['class']}")
    add_centered(doc, f"{lab['total']}: {data['total_hours']} {lab['hours']}, {lab['per_week']}: {data['hours_per_week']} {lab['hours']}")
    add_centered(doc, f"{lab['year']}: 2026–2027")

    p = doc.add_paragraph()
    p.add_run(f"{lab['school_week']}: ").bold = True
    p.add_run(school_week_text(data.get("school_week", "5-day"), str(data["language"])))
    apply_font_to_paragraph(p)

    p = doc.add_paragraph()
    p.add_run(f"Источник календаря / Күнтізбе дереккөзі / Calendar source: ").bold = True
    p.add_run(str(data["calendar_source"]))
    apply_font_to_paragraph(p)

    table = add_plan_table(doc, lab)

    for index, quarter in enumerate(data["quarters"], 1):
        qrow = table.add_row()
        merged = qrow.cells[0].merge(qrow.cells[5])
        qname = quarter.get("name") or f"{index} {lab['quarter']}"
        set_cell_text(merged, f"{qname} — {quarter['weeks']} {lab['weeks']}", True, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(qrow.cells[6], "", True, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(merged, "FFFFFF")
        set_cell_shading(qrow.cells[6], "FFFFFF")
        for paragraph in merged.paragraphs + qrow.cells[6].paragraphs:
            paragraph.paragraph_format.keep_with_next = True
        prevent_row_split(qrow)
        for lesson in quarter.get("lessons", []):
            row = table.add_row()
            assessment_type = str(lesson.get("assessment_type", "")).strip()
            note = str(lesson.get("note", "")).strip()
            displayed_note = " — ".join(value for value in (assessment_type, note) if value)
            values = [
                lesson.get("number", ""), lesson.get("section", ""), lesson.get("topic", ""),
                lesson.get("learning_objectives", ""), lesson.get("hours", lesson.get("hours_total", "")),
                lesson.get("date", ""), displayed_note
            ]
            for i, value in enumerate(values):
                align = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
                set_cell_text(row.cells[i], value, False, align)
            prevent_row_split(row)

    conflicts = data.get("source_conflicts", [])
    if conflicts:
        heading = doc.add_paragraph()
        heading.add_run("Выявленные противоречия официального источника").bold = True
        apply_font_to_paragraph(heading)
        conflict_table = doc.add_table(rows=1, cols=4)
        conflict_table.style = "Table Grid"
        headings = ["Место", "Тип", "Исходный фрагмент", "Решение"]
        for index, value in enumerate(headings):
            set_cell_text(conflict_table.rows[0].cells[index], value, True, WD_ALIGN_PARAGRAPH.CENTER)
        set_repeat_table_header(conflict_table.rows[0])
        for conflict in conflicts:
            row = conflict_table.add_row()
            values = [
                conflict["location"], conflict["type"],
                conflict["source_text"], conflict["resolution"],
            ]
            for index, value in enumerate(values):
                set_cell_text(row.cells[index], value)
            prevent_row_split(row)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    audit_docx_structure(
        output, int(data["total_hours"]), bool(conflicts), expected_critical_unicode(data)
    )


def audit_docx_structure(
    output: Path,
    total_hours: int,
    has_conflicts: bool,
    expected_unicode: set[str] | None = None,
) -> None:
    """Reject fragmented KTP tables and manual pagination after saving."""
    document = Document(output)
    expected_table_count = 2 if has_conflicts else 1
    if len(document.tables) != expected_table_count:
        raise ValueError(
            f"DOCX must contain one main KTP table"
            f"{' and one conflict appendix table' if has_conflicts else ''}; "
            f"found {len(document.tables)} tables"
        )
    main_table = document.tables[0]
    expected_rows = 1 + 4 + total_hours
    if len(main_table.rows) != expected_rows:
        raise ValueError(
            f"Main KTP table must contain {expected_rows} rows, found {len(main_table.rows)}"
        )
    if main_table._tbl.tblGrid is None:
        raise ValueError("Main KTP table must contain w:tblGrid")
    header_rows = main_table._tbl.xpath("./w:tr[w:trPr/w:tblHeader]")
    if len(header_rows) != 1 or header_rows[0] is not main_table.rows[0]._tr:
        raise ValueError("Main KTP table must have exactly one physical header row marked w:tblHeader")
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml")
    decoded_xml = xml.decode("utf-8", "strict")
    if "\ufffd" in decoded_xml:
        raise ValueError("DOCX contains the Unicode replacement character U+FFFD")
    round_trip_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    for char in expected_unicode or set():
        if char not in decoded_xml:
            raise ValueError(f"DOCX did not preserve required Unicode character {char!r}")
        if char not in round_trip_text:
            raise ValueError(f"Word text round-trip did not preserve Unicode character {char!r}")
    for r_fonts in re.findall(rb"<w:rFonts\b[^>]*/>", xml):
        for channel in (b"ascii", b"hAnsi", b"eastAsia", b"cs"):
            match = re.search(rb"w:" + channel + rb"=[\"']([^\"']+)[\"']", r_fonts)
            if match is None or match.group(1).decode("utf-8") != FONT:
                raise ValueError(
                    "Every text run must set Times New Roman in ascii, hAnsi, eastAsia, and cs"
                )
    if re.search(rb'<w:br\b[^>]*w:type=["\']page["\']', xml) or b"<w:pageBreakBefore" in xml:
        raise ValueError("Manual page breaks and pageBreakBefore are prohibited in the KTP DOCX")
    fills = re.findall(rb'<w:shd\b[^>]*w:fill=["\']([^"\']+)["\']', xml)
    text_colors = re.findall(rb'<w:color\b[^>]*w:val=["\']([^"\']+)["\']', xml)
    if any(value.upper() not in {b"FFFFFF", b"AUTO"} for value in fills):
        raise ValueError("Colored shading is prohibited in the official monochrome KTP DOCX")
    if any(value.upper() not in {b"000000", b"AUTO"} for value in text_colors):
        raise ValueError("Colored text is prohibited in the official monochrome KTP DOCX")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output")
    parser.add_argument("--output-dir")
    args = parser.parse_args()
    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    if args.output:
        output = Path(args.output)
    else:
        output_dir = Path(args.output_dir or ".")
        output = output_dir / default_output_name(data)
    build(data, output)
    print(output)


if __name__ == "__main__":
    main()
