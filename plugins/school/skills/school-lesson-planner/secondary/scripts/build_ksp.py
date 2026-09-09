#!/usr/bin/env python3
"""Build and structurally audit a Kazakhstan KSP DOCX from validated JSON."""

from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATES = {
    "kk": SKILL_DIR / "assets" / "ksp-template-kk.docx",
    "ru": SKILL_DIR / "assets" / "ksp-template-ru.docx",
    "en": SKILL_DIR / "assets" / "ksp-template-en.docx",
}
LANGUAGE_ALIASES = {
    "kk": "kk", "kz": "kk", "kazakh": "kk", "қазақ": "kk", "қазақша": "kk",
    "ru": "ru", "russian": "ru", "русский": "ru", "рус": "ru",
    "en": "en", "english": "en", "английский": "en",
}
LABELS = {
    "kk": {
        "organization": "Білім беру ұйымының атауы: ",
        "minute": "мин",
        "appendix": "Әдістемелік қосымша",
        "goal_analysis": "Оқу мақсаттарын әдістемелік талдау",
        "knowledge_skills_analysis": "Оқу мақсаты мен сабақ тақырыбы арқылы дамитын білім мен дағдылар",
        "methodology_application": "Таңдалған әдістемені сабақта қолдану",
        "reflection_use": "Алдыңғы сабақ рефлексиясын пайдалану",
        "model_rationale": "Сабақ моделін таңдау негіздемесі",
        "term_explanations": "Терминдерге қысқаша түсініктеме",
        "differentiation": "Саралау және инклюзивті дизайн",
        "formative_assessment": "Қалыптастырушы бағалау логикасы",
        "alternatives": "Балама және резервтік тапсырмалар",
        "external_resources": "Дереккөздер",
        "professional_notice": "Маңызды ескерту",
        "descriptor": "Дескриптор",
        "feedback": "Кері байланыс",
        "professional_notice_text": "Құрметті әріптес, назар аударыңыз, сабақ жоспары және оның әдістемелік қосымшасы ұсыныс ретінде құрастырылды. Құрастырылған оқу материалдары педагог тарапынан міндетті түрде тексеруді талап етеді. Жасанды интеллект педагогикалық шешімді алмастыра алмайтындығын еске саламыз.",
    },
    "ru": {
        "organization": "Наименование организации образования: ",
        "minute": "мин",
        "appendix": "Методическое приложение",
        "goal_analysis": "Методический анализ целей обучения",
        "knowledge_skills_analysis": "Знания и навыки, реализуемые через цель обучения и тему урока",
        "methodology_application": "Применение выбранной методики в уроке",
        "reflection_use": "Использование рефлексии предыдущего урока",
        "model_rationale": "Обоснование модели урока",
        "term_explanations": "Краткое пояснение терминов",
        "differentiation": "Дифференциация и инклюзивный дизайн",
        "formative_assessment": "Логика формативного оценивания",
        "alternatives": "Альтернативные и резервные задания",
        "external_resources": "Источники",
        "professional_notice": "Важное примечание",
        "descriptor": "Дескриптор",
        "feedback": "Обратная связь",
        "professional_notice_text": "Уважаемый коллега, обратите внимание: план урока и его методическое приложение составлены в качестве рекомендации. Разработанные учебные материалы требуют обязательной проверки педагогом. Искусственный интеллект не может заменить педагогическое решение.",
    },
    "en": {
        "organization": "Name of educational organization: ",
        "minute": "min",
        "appendix": "Methodological appendix",
        "goal_analysis": "Methodological analysis of learning objectives",
        "knowledge_skills_analysis": "Knowledge and skills developed through the learning objective and lesson topic",
        "methodology_application": "Application of the selected methodology",
        "reflection_use": "Use of previous-lesson reflection",
        "model_rationale": "Rationale for the lesson model",
        "term_explanations": "Brief explanation of terms",
        "differentiation": "Differentiation and inclusive design",
        "formative_assessment": "Formative-assessment logic",
        "alternatives": "Alternative and reserve tasks",
        "external_resources": "Sources",
        "professional_notice": "Important notice",
        "descriptor": "Descriptor",
        "feedback": "Feedback",
        "professional_notice_text": "Dear colleague, please note: the lesson plan and its methodological appendix have been prepared as recommendations. The developed learning materials must be reviewed by the teacher. Artificial intelligence cannot replace professional pedagogical judgment.",
    },
}
APPENDIX_ORDER = (
    "goal_analysis", "knowledge_skills_analysis", "methodology_application",
    "reflection_use", "model_rationale", "differentiation",
    "formative_assessment", "alternatives", "external_resources",
)
APPENDIX_REQUIRED = {
    "goal_analysis", "knowledge_skills_analysis", "methodology_application",
    "model_rationale", "differentiation",
    "formative_assessment", "alternatives", "audit_summary",
}
LESSON_DURATION_MINUTES = 40
REQUIRED_ROOT = {
    "intake_verification", "instruction_language", "language", "lesson_count", "subject",
    "grade", "section", "topic", "class_size", "teacher_experience",
    "qualification_category", "class_characteristics",
    "learning_objectives", "lesson_objectives", "assessment_criteria",
    "stages", "methodological_appendix",
}
REQUIRED_INTAKE_FIELDS = {
    "subject", "grade", "instruction_language", "section", "topic", "learning_objectives",
    "lesson_count", "class_size", "teacher_experience",
    "qualification_category", "class_characteristics",
}
REQUIRED_STAGE = {
    "name", "minutes", "teacher_actions", "learner_actions", "assessment", "resources",
}
REQUIRED_STAGE_V3 = {
    "name", "minutes", "activity_type", "methods", "teacher_actions", "tasks", "resources",
}
REQUIRED_TASK = {
    "canonical_task_id", "task_type", "objective_refs", "instruction", "learner_action",
    "expected_product", "descriptor", "feedback",
}
REQUIRED_METHOD = {"method_id", "type", "name"}
ACTIVITY_TYPES = {"organization", "learning"}
TASK_TYPES = {"oral", "written", "practical", "laboratory", "group", "reflection", "homework"}
METHOD_TYPES = {"primary", "supporting", "technique"}
sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "scripts"))
from lesson_table_policy import contains_framework
from scored_tasks import scoring
WORK_FORMS = {
    "ru": {"individual": "Индивидуальная работа", "pair": "Парная работа", "group": "Групповая работа", "whole_class": "Фронтальная работа"},
    "kk": {"individual": "Жеке жұмыс", "pair": "Жұптық жұмыс", "group": "Топтық жұмыс", "whole_class": "Ұжымдық жұмыс"},
    "en": {"individual": "Individual work", "pair": "Pair work", "group": "Group work", "whole_class": "Whole-class work"},
}


def scored_heading(language: str, total: int) -> str:
    if language == "ru":
        unit = "баллов" if 11 <= total % 100 <= 14 else ("балл" if total % 10 == 1 else "балла" if 2 <= total % 10 <= 4 else "баллов")
        return f"Дескрипторы — {total} {unit}:"
    if language == "kk":
        return f"Дескрипторлар — {total} балл:"
    return f"Descriptors — {total} {'point' if total == 1 else 'points'}:"


SUPPORT_LABEL = {"ru": "Поддержка: ", "kk": "Қолдау: ", "en": "Support: "}
TASK_LABEL = {"ru": "Задание", "kk": "Тапсырма", "en": "Task"}
METHOD_PREFIX = {
    "kk": "Әдіс-тәсіл: ",
    "ru": "Метод/приём: ",
    "en": "Method/technique: ",
}
OBJECTIVE_CODE = re.compile(r"(?<!\d)\d{1,2}(?:\.\d+){2,}(?!\d)")
BANNED = (
    "жетістік критерийлері", "критерии успеха", "success criteria",
    "күтілетін нәтиже", "ожидаемый результат", "expected result",
)
OLD_ENDING_TERMS = ("шығу билеті", "выходной билет")
EMPHASIZED_ITEM = re.compile(r"^\*\*\*(?P<text>.+?)\*\*\*$", re.DOTALL)
NON_DESCRIPTORS = {
    "ауызша кері байланыс", "устная обратная связь", "oral feedback",
    "критерийлер бойынша тексеру", "проверка по критериям", "criteria check",
    "өзін-өзі бағалау", "самооценивание", "self-assessment",
}
DEPRECATED_ASSESSMENT_NOTES = (
    "ескерту: «бағалау критерийлері»",
    "примечание: «критерии оценивания»",
    "note: “assessment criteria”",
    'note: "assessment criteria"',
)
PRACTICAL_SUBJECT = re.compile(
    r"(?:хими|chemistr|химия|биологи|biology|биология|физик|physics|физика|"
    r"жаратылыстану|естествозн|natural\s+science|көркем\s+еңбек|"
    r"художественн(?:ый|ого)\s+труд|arts?\s*(?:and|&)\s*crafts?)",
    re.IGNORECASE,
)
KNOWLEDGE_SKILL_PREFIXES = {
    "kk": (
        "ПӘНДІК МАЗМҰН:",
        "ФАКТІЛІК БІЛІМ:",
        "ПӘНДІК ДАҒДЫЛАР:",
        "ТАНЫМДЫҚ ДАҒДЫЛАР:",
    ),
    "ru": (
        "ПРЕДМЕТНОЕ СОДЕРЖАНИЕ:",
        "ФАКТИЧЕСКИЕ ЗНАНИЯ:",
        "ПРЕДМЕТНЫЕ НАВЫКИ:",
        "ПОЗНАВАТЕЛЬНЫЕ НАВЫКИ:",
    ),
    "en": (
        "SUBJECT CONTENT:",
        "FACTUAL KNOWLEDGE:",
        "SUBJECT-SPECIFIC SKILLS:",
        "COGNITIVE SKILLS:",
    ),
}
REMOVED_KNOWLEDGE_SKILL_TERMS = (
    "тұжырымдамалық білім", "процедуралық білім", "метатанымдық білім",
    "метакогнитивтік білім", "қосымша дағдылар",
    "концептуальные знания", "концептуальное знание",
    "процедурные знания", "процедурное знание",
    "метакогнитивные знания", "метакогнитивное знание",
    "дополнительные навыки", "conceptual knowledge", "procedural knowledge",
    "metacognitive knowledge", "additional skills",
)
METHOD_TERM_PATTERNS = {
    "UbD": re.compile(r"(?<![\w-])(?:ubd|understanding by design)(?![\w-])", re.IGNORECASE),
    "UDL": re.compile(r"(?<![\w-])(?:udl|universal design for learning)(?![\w-])", re.IGNORECASE),
    "Visible Learning": re.compile(r"(?<![\w-])visible learning(?![\w-])", re.IGNORECASE),
    "Explicit Instruction": re.compile(r"(?<![\w-])explicit instruction(?![\w-])", re.IGNORECASE),
    "Wiliam": re.compile(r"(?<![\w-])wiliam(?![\w-])", re.IGNORECASE),
    "Marzano": re.compile(r"(?<![\w-])marzano(?![\w-])", re.IGNORECASE),
    "Tomlinson": re.compile(r"(?<![\w-])tomlinson(?![\w-])", re.IGNORECASE),
    "Archer-Hughes": re.compile(r"(?<![\w-])archer[–—-]hughes(?![\w-])", re.IGNORECASE),
    "Hattie": re.compile(r"(?<![\w-])hattie(?![\w-])", re.IGNORECASE),
    "Agarwal-Bain": re.compile(r"(?<![\w-])agarwal[–—-]bain(?![\w-])", re.IGNORECASE),
    "Rosenshine": re.compile(r"(?<![\w-])rosenshine(?:’s|'s)?(?![\w-])", re.IGNORECASE),
    "Visible Thinking": re.compile(r"(?<![\w-])visible thinking(?![\w-])", re.IGNORECASE),
    "Dialogic Teaching": re.compile(r"(?<![\w-])dialogic teaching(?![\w-])", re.IGNORECASE),
    "Project Based Learning": re.compile(r"(?<![\w-])(?:project based learning|pbl)(?![\w-])", re.IGNORECASE),
    "Design Thinking": re.compile(r"(?<![\w-])design thinking(?![\w-])", re.IGNORECASE),
    "Inquiry-based Learning": re.compile(r"(?<![\w-])inquiry[- ]based learning(?![\w-])", re.IGNORECASE),
    "Problem-based Learning": re.compile(r"(?<![\w-])problem[- ]based learning(?![\w-])", re.IGNORECASE),
}
TERM_STATUS_MARKERS = {
    "kk": "әдістемелік ұсыныс",
    "ru": "методическая рекомендация",
    "en": "methodological recommendation",
}
KK_SOURCE_STATUS_GOAL = "Оқу мақсаттары — мұғалім ұсынған үлгіде граматикалық ерекшеліктері сақталды."
KK_SOURCE_STATUS_EXPLICIT = "Anita L. Archer және Charles A. Hughes еңбегіндегі Explicit Instruction қағидалары — осы сабаққа бейімделген әдістемелік ұсыныс."
KK_SOURCE_STATUS_FORM = "ҚМЖ нысаны — плагиннің 15.08.2026 күні жаңартылған №130 бұйрық жөніндегі нормативтік анықтамасына сүйеніп рәсімделді."
RU_SOURCE_STATUS_GOAL = "Цели обучения использованы в формулировке, предоставленной учителем, с сохранением грамматических особенностей."
RU_SOURCE_STATUS_EXPLICIT = "Принципы Explicit Instruction из работы Anita L. Archer и Charles A. Hughes адаптированы для данного урока и представлены как методическая рекомендация."
RU_SOURCE_STATUS_FORM = "Форма КСП оформлена на основании нормативной справки плагина о приказе №130, обновлённой 15.08.2026."
EN_SOURCE_STATUS_GOAL = "The learning objectives are presented in the wording provided by the teacher, with their grammatical features preserved."
EN_SOURCE_STATUS_EXPLICIT = "The Explicit Instruction principles described by Anita L. Archer and Charles A. Hughes have been adapted for this lesson as a methodological recommendation."
EN_SOURCE_STATUS_FORM = "The lesson-plan form was prepared using the plugin's regulatory reference for Order No. 130, updated on August 15, 2026."
LESSON_WIDTHS_DXA = (1656, 2880, 2448, 1728, 1584)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


class KSPError(ValueError):
    pass


def nonempty(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (list, tuple)):
        return bool(value) and all(nonempty(item) for item in value)
    return True


def content_items(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value.strip()] if value.strip() else []
    if isinstance(value, list):
        result = []
        for item in value:
            if isinstance(item, str) and item.strip():
                result.append(item.strip())
            else:
                raise KSPError("Text lists may contain only non-empty strings.")
        return result
    raise KSPError("Text content must be a string or a list of strings.")


def bullet_items(value: Any, field_name: str) -> list[str]:
    """Require one clean JSON array item for every Word bullet."""
    if not isinstance(value, list):
        raise KSPError(
            f"{field_name} must be a JSON array with one separate item per bullet."
        )
    items = content_items(value)
    for item in items:
        if "\n" in item or "\r" in item:
            raise KSPError(
                f"{field_name} items must not contain line breaks; "
                "put every objective or criterion in a separate array item."
            )
        if re.match(r"^\s*(?:[•◦▪‣⁃*-]|\d+[.)])\s+", item):
            raise KSPError(
                f"{field_name} items must not contain typed bullet or number markers; "
                "the DOCX builder adds real Word bullets."
            )
    return items


def normalized_language(value: Any) -> str:
    key = str(value or "").strip().casefold()
    if key not in LANGUAGE_ALIASES:
        raise KSPError("language must be kk, ru, or en.")
    return LANGUAGE_ALIASES[key]


def normalized_field_name(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value).replace("_", " ").replace("-", " ").casefold()).strip()


def normalized_text(value: Any) -> str:
    return re.sub(r"[^\w]+", " ", str(value).casefold(), flags=re.UNICODE).strip()


def validate_task(
    task: Any,
    *,
    stage_index: int,
    seen_task_ids: set[str],
    descriptor_owners: dict[str, list[str]],
) -> dict[str, Any]:
    if not isinstance(task, dict):
        raise KSPError(f"Stage {stage_index} task must be an object.")
    missing = sorted(key for key in REQUIRED_TASK if key not in task or not nonempty(task[key]))
    if missing:
        raise KSPError(f"Stage {stage_index} task is missing: " + ", ".join(missing))
    result = deepcopy(task)
    task_id = str(result["canonical_task_id"]).strip()
    if task_id in seen_task_ids:
        raise KSPError(f"Duplicate canonical_task_id: {task_id}")
    seen_task_ids.add(task_id)
    result["canonical_task_id"] = task_id
    task_type = str(result["task_type"]).strip()
    if task_type not in TASK_TYPES:
        raise KSPError(f"Task {task_id} has unsupported task_type: {task_type}")
    result["task_type"] = task_type
    refs = content_items(result["objective_refs"])
    result["objective_refs"] = refs
    for key in ("instruction", "learner_action", "expected_product", "descriptor", "feedback"):
        if not isinstance(result[key], str) or not result[key].strip():
            raise KSPError(f"Task {task_id} field {key} must be non-empty text.")
        result[key] = result[key].strip()
    if normalized_text(result["descriptor"]) == normalized_text(result["instruction"]):
        raise KSPError(f"Descriptor repeats the instruction for {task_id}.")
    if normalized_text(result["descriptor"]) in {normalized_text(value) for value in NON_DESCRIPTORS}:
        raise KSPError(f"Task {task_id} uses feedback or assessment mode instead of a descriptor.")
    descriptor_owners.setdefault(normalized_text(result["descriptor"]), []).append(task_id)
    scores = result.get("descriptor_scores")
    if not isinstance(scores, list) or not scores:
        raise KSPError(f"Task {task_id} requires descriptor_scores with explicit points.")
    for item in scores:
        if (not isinstance(item, dict) or not isinstance(item.get("text"), str)
                or not item["text"].strip() or type(item.get("points")) is not int
                or item["points"] <= 0):
            raise KSPError(f"Task {task_id}: each descriptor score needs text and positive integer points.")
    total = sum(item["points"] for item in scores)
    if "total_points" in result and (type(result["total_points"]) is not int or result["total_points"] != total):
        raise KSPError(f"Task {task_id}: total_points must equal the descriptor sum.")
    result["total_points"] = total
    if "support" in result and not isinstance(result["support"], str):
        raise KSPError(f"Task {task_id}: support must be text.")
    result.update(scoring(result))
    return result


def validate_method(method: Any, *, stage_index: int, seen_method_ids: set[str]) -> dict[str, str]:
    if not isinstance(method, dict):
        raise KSPError(f"Stage {stage_index} method must be an object.")
    missing = sorted(key for key in REQUIRED_METHOD if key not in method or not nonempty(method[key]))
    if missing:
        raise KSPError(f"Stage {stage_index} method is missing: " + ", ".join(missing))
    result = {key: str(method[key]).strip() for key in REQUIRED_METHOD}
    if result["method_id"] in seen_method_ids:
        raise KSPError(f"Duplicate method_id: {result['method_id']}")
    seen_method_ids.add(result["method_id"])
    if result["type"] not in METHOD_TYPES:
        raise KSPError(f"Method {result['method_id']} has unsupported type: {result['type']}")
    if "***" in result["name"] or ":" in result["name"][:24]:
        raise KSPError(f"Method {result['method_id']} must contain only the method name, without markup or prefix.")
    if contains_framework(result["name"]):
        raise KSPError("Put methodology frameworks in the methodological appendix; stage methods must name practical techniques.")
    if any(result["name"].casefold() == label.casefold() for labels in WORK_FORMS.values() for label in labels.values()):
        raise KSPError("A work form is not a practical technique.")
    return result


def find_banned_field_names(value: Any) -> list[str]:
    found: list[str] = []
    if isinstance(value, dict):
        for key, child in value.items():
            normalized = normalized_field_name(key)
            if any(term in normalized for term in BANNED):
                found.append(str(key))
            found.extend(find_banned_field_names(child))
    elif isinstance(value, list):
        for child in value:
            found.extend(find_banned_field_names(child))
    return found


def validate_intake(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise KSPError("intake_verification must be an object.")
    if value.get("confirmed_by_user") is not True:
        raise KSPError("The mandatory teacher interview must be explicitly confirmed by the user.")
    source_summary = value.get("source_summary")
    if not isinstance(source_summary, str) or not source_summary.strip():
        raise KSPError("intake_verification.source_summary must identify the user-provided source.")
    fields = value.get("confirmed_fields")
    if not isinstance(fields, list) or not all(isinstance(item, str) for item in fields):
        raise KSPError("intake_verification.confirmed_fields must be a list of field names.")
    normalized = {normalized_field_name(item).replace(" ", "_") for item in fields}
    missing = sorted(REQUIRED_INTAKE_FIELDS - normalized)
    if missing:
        raise KSPError("The mandatory teacher interview is missing confirmed fields: " + ", ".join(missing))
    return value


def validate_input(raw: Any) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise KSPError("The JSON root must be an object.")
    duration = raw.get("lesson_duration_minutes", LESSON_DURATION_MINUTES)
    if isinstance(duration, bool) or not isinstance(duration, int) or duration <= 0:
        raise KSPError("lesson_duration_minutes must be a positive integer when explicitly supplied.")
    missing = sorted(key for key in REQUIRED_ROOT if key not in raw or not nonempty(raw[key]))
    if missing:
        raise KSPError("Missing required fields: " + ", ".join(missing))

    data = deepcopy(raw)
    schema_version = str(data.get("schema_version", "2.0")).strip()
    if schema_version not in {"2.0", "3.0"}:
        raise KSPError("schema_version must be 2.0 or 3.0.")
    data["schema_version"] = schema_version
    data["validation_status"] = "strict" if schema_version == "3.0" else "legacy_unverified"
    data["validation_warnings"] = []
    data["lesson_duration_minutes"] = duration
    data["intake_verification"] = validate_intake(data["intake_verification"])
    data["language"] = normalized_language(data["language"])
    if PRACTICAL_SUBJECT.search(str(data["subject"])):
        if "practical_resources" not in data or not nonempty(data["practical_resources"]):
            raise KSPError(
                "practical_resources is required for chemistry, biology, physics, "
                "natural science, and arts/crafts."
            )
        confirmed = {
            normalized_field_name(item).replace(" ", "_")
            for item in data["intake_verification"]["confirmed_fields"]
        }
        if "practical_resources" not in confirmed:
            raise KSPError(
                "The teacher must explicitly confirm practical_resources for this subject."
            )
        if not isinstance(data["practical_resources"], str) or not data["practical_resources"].strip():
            raise KSPError("practical_resources must be non-empty teacher-provided text.")
    lesson_count = data["lesson_count"]
    if isinstance(lesson_count, bool) or not isinstance(lesson_count, int) or lesson_count <= 0:
        raise KSPError("lesson_count must be a positive integer confirmed by the teacher.")

    for key in (
        "subject", "grade", "section", "topic", "teacher_experience",
        "qualification_category", "class_characteristics",
    ):
        if not isinstance(data[key], (str, int)) or not str(data[key]).strip():
            raise KSPError(f"{key} must be non-empty text or a number.")
    class_size = data["class_size"]
    if isinstance(class_size, bool) or not isinstance(class_size, int) or class_size <= 0:
        raise KSPError("class_size must be a positive integer confirmed by the teacher.")

    objectives = content_items(data["learning_objectives"])
    for objective in objectives:
        if not OBJECTIVE_CODE.search(objective):
            raise KSPError(
                "Each learning objective must include its supplied curriculum code; "
                f"no code was found in: {objective!r}"
            )
    data["learning_objectives"] = objectives
    supplied_objective_codes = {
        match.group(0)
        for objective in objectives
        for match in OBJECTIVE_CODE.finditer(objective)
    }
    data["lesson_objectives"] = bullet_items(data["lesson_objectives"], "lesson_objectives")
    data["assessment_criteria"] = bullet_items(data["assessment_criteria"], "assessment_criteria")

    if not isinstance(data["stages"], list) or not data["stages"]:
        raise KSPError("stages must be a non-empty list.")
    total = 0
    seen_task_ids: set[str] = set()
    seen_method_ids: set[str] = set()
    descriptor_owners: dict[str, list[str]] = {}
    primary_method_seen = False
    for index, stage in enumerate(data["stages"], start=1):
        if not isinstance(stage, dict):
            raise KSPError(f"Stage {index} must be an object.")
        if schema_version == "3.0":
            missing_stage = sorted(key for key in REQUIRED_STAGE_V3 if key not in stage)
            if missing_stage:
                raise KSPError(f"Stage {index} is missing: " + ", ".join(missing_stage))
            for key in ("name", "minutes", "teacher_actions", "resources"):
                if not nonempty(stage[key]):
                    raise KSPError(f"Stage {index} is missing: {key}")
            activity_type = str(stage["activity_type"]).strip()
            if activity_type not in ACTIVITY_TYPES:
                raise KSPError(f"Stage {index} has unsupported activity_type: {activity_type}")
            stage["activity_type"] = activity_type
            forms = stage.get("work_forms", [])
            if (not isinstance(forms, list)
                    or any(not isinstance(form, str) or form not in WORK_FORMS["ru"] for form in forms)
                    or (activity_type == "learning" and not forms)):
                raise KSPError(f"Stage {index} requires valid work_forms.")
            stage["teacher_actions"] = content_items(stage["teacher_actions"])
            bindings = stage.get("action_work_forms")
            if bindings is None and len(forms) == 1:
                bindings = [forms[:] for _ in stage["teacher_actions"]]
            if activity_type == "learning" and (not isinstance(bindings, list)
                    or len(bindings) != len(stage["teacher_actions"])
                    or any(not isinstance(group, list) or not group or any(form not in forms for form in group) for group in bindings)
                    or set(form for group in bindings for form in group) != set(forms)):
                raise KSPError("action_work_forms must link each action to its explicit work forms; multiple forms cannot be inferred.")
            stage["action_work_forms"] = bindings or [[] for _ in stage["teacher_actions"]]
            if contains_framework(" ".join(stage["teacher_actions"])):
                raise KSPError("Move methodology explanations from teacher actions to the methodological appendix.")
            if any("***" in action for action in stage["teacher_actions"]):
                raise KSPError(f"Stage {index} teacher_actions must not contain Markdown emphasis markers.")
            if not isinstance(stage["methods"], list):
                raise KSPError(f"Stage {index} methods must be a list.")
            stage["methods"] = [
                validate_method(method, stage_index=index, seen_method_ids=seen_method_ids)
                for method in stage["methods"]
            ]
            if activity_type == "learning" and not stage["methods"]:
                raise KSPError("Every learning stage requires a concrete practical technique.")
            primary_method_seen = primary_method_seen or any(
                method["type"] == "primary" for method in stage["methods"]
            )
            if not isinstance(stage["tasks"], list):
                raise KSPError(f"Stage {index} tasks must be a list.")
            if activity_type == "learning" and not stage["tasks"]:
                raise KSPError(f"Stage {index} is a learning activity and requires at least one task.")
            if activity_type == "organization" and stage["tasks"]:
                raise KSPError(f"Stage {index} is organizational and must not contain learning tasks.")
            stage["tasks"] = [
                validate_task(
                    task,
                    stage_index=index,
                    seen_task_ids=seen_task_ids,
                    descriptor_owners=descriptor_owners,
                )
                for task in stage["tasks"]
            ]
            for task in stage["tasks"]:
                unknown_refs = sorted(set(task["objective_refs"]) - supplied_objective_codes)
                if unknown_refs:
                    raise KSPError(
                        f"Task {task['canonical_task_id']} references objectives not supplied by the teacher: "
                        + ", ".join(unknown_refs)
                    )
        else:
            missing_stage = sorted(key for key in REQUIRED_STAGE if key not in stage or not nonempty(stage[key]))
            if missing_stage:
                raise KSPError(f"Stage {index} is missing: " + ", ".join(missing_stage))
        minutes = stage["minutes"]
        if isinstance(minutes, bool) or not isinstance(minutes, int) or minutes <= 0:
            raise KSPError(f"Stage {index} minutes must be a positive integer.")
        total += minutes
        values_to_check = (
            stage["teacher_actions"]
            if schema_version == "3.0"
            else [stage[key] for key in REQUIRED_STAGE - {"minutes"}]
        )
        for value in values_to_check:
            items = content_items(value) if not isinstance(value, str) else [value]
            for item in items:
                if any(term in item.casefold() for term in OLD_ENDING_TERMS):
                    raise KSPError(
                        "Use the localized final-stage term Reflection/Рефлексия; "
                        "use reflection and actionable feedback; do not use former labels for the final stage."
                    )
    if schema_version == "3.0":
        if not primary_method_seen:
            raise KSPError("Schema 3.0 requires at least one primary method in the stage where it is used.")
        duplicates = [owners for owners in descriptor_owners.values() if len(owners) > 1]
        for owners in duplicates:
            data["validation_warnings"].append(
                "Repeated descriptor text is mapped separately to: " + ", ".join(owners)
            )
    if total != data["lesson_duration_minutes"]:
        raise KSPError(
            f"Stage minutes total {total}, but this lesson must total {data['lesson_duration_minutes']}."
        )

    instruction_language = str(data["instruction_language"]).strip()
    if instruction_language not in {"ru", "kk", "en"}:
        raise KSPError("instruction_language must be ru, kk, or en.")
    if data["language"] != instruction_language and data.get("document_language_explicit") is not True:
        raise KSPError("Document language must match instruction_language unless the teacher explicitly requested an override.")

    appendix = data["methodological_appendix"]
    if not isinstance(appendix, dict):
        raise KSPError("methodological_appendix must be an object.")
    missing_appendix = sorted(key for key in APPENDIX_REQUIRED if key not in appendix or not nonempty(appendix[key]))
    if missing_appendix:
        raise KSPError("The methodological appendix is missing: " + ", ".join(missing_appendix))
    for key in APPENDIX_ORDER:
        if key in appendix and nonempty(appendix[key]):
            content_items(appendix[key])

    knowledge_items = content_items(appendix["knowledge_skills_analysis"])
    required_prefixes = KNOWLEDGE_SKILL_PREFIXES[data["language"]]
    if len(knowledge_items) != len(required_prefixes):
        raise KSPError(
            "knowledge_skills_analysis must contain exactly four localized items: "
            "subject content, factual knowledge, subject-specific skills, and cognitive skills."
        )
    for index, (item, prefix) in enumerate(zip(knowledge_items, required_prefixes), start=1):
        if not item.casefold().startswith(prefix.casefold()):
            raise KSPError(
                f"knowledge_skills_analysis item {index} must start with {prefix!r}."
            )
        removed = [term for term in REMOVED_KNOWLEDGE_SKILL_TERMS if term in item.casefold()]
        if removed:
            raise KSPError(
                "knowledge_skills_analysis contains a removed category: "
                + ", ".join(removed)
            )

    terminology_source = "\n".join(
        item
        for key, value in appendix.items()
        if key not in {"term_explanations", "audit_summary"}
        for item in content_items(value)
    )
    used_methods = [
        name for name, pattern in METHOD_TERM_PATTERNS.items()
        if pattern.search(terminology_source)
    ]
    if used_methods:
        if "term_explanations" not in appendix or not nonempty(appendix["term_explanations"]):
            raise KSPError(
                "term_explanations is required when a professional method name or acronym is used: "
                + ", ".join(used_methods)
            )
        explanations = "\n".join(content_items(appendix["term_explanations"]))
        missing_terms = [
            name for name in used_methods
            if not METHOD_TERM_PATTERNS[name].search(explanations)
        ]
        if missing_terms:
            raise KSPError(
                "term_explanations must explain each named method: "
                + ", ".join(missing_terms)
            )
        status_marker = TERM_STATUS_MARKERS[data["language"]]
        if status_marker.casefold() not in explanations.casefold():
            raise KSPError(
                "term_explanations must identify the pedagogical model as a "
                f"{status_marker}."
            )

    fixed_sources = {
        "kk": (KK_SOURCE_STATUS_GOAL, KK_SOURCE_STATUS_EXPLICIT, KK_SOURCE_STATUS_FORM),
        "ru": (RU_SOURCE_STATUS_GOAL, RU_SOURCE_STATUS_EXPLICIT, RU_SOURCE_STATUS_FORM),
        "en": (EN_SOURCE_STATUS_GOAL, EN_SOURCE_STATUS_EXPLICIT, EN_SOURCE_STATUS_FORM),
    }
    goal_source, explicit_source, form_source = fixed_sources[data["language"]]
    source_status = [goal_source]
    if "Explicit Instruction" in used_methods:
        source_status.append(explicit_source)
    elif used_methods:
        method_name = used_methods[0]
        generic_method_source = {
            "kk": f"{method_name} оқыту әдістемесі — осы сабаққа бейімделген әдістемелік ұсыныс.",
            "ru": f"Методика обучения {method_name} адаптирована для данного урока и представлена как методическая рекомендация.",
            "en": f"The {method_name} teaching methodology has been adapted for this lesson as a methodological recommendation.",
        }
        source_status.append(generic_method_source[data["language"]])
    else:
        generic_source = {
            "kk": "Сабақта көрсетілген оқыту әдістемесі — осы сабаққа бейімделген әдістемелік ұсыныс.",
            "ru": "Указанная в уроке методика обучения адаптирована для данного урока и представлена как методическая рекомендация.",
            "en": "The teaching methodology described in the lesson has been adapted as a methodological recommendation.",
        }
        source_status.append(generic_source[data["language"]])
    source_status.append(form_source)
    appendix["external_resources"] = source_status

    banned = find_banned_field_names(data)
    if banned:
        raise KSPError("Banned field name found: " + ", ".join(dict.fromkeys(banned)))
    return data


def clear_paragraph(paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def remove_deprecated_template_notes(doc) -> None:
    """Remove the retired teacher-facing note retained in older bundled templates."""
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip().casefold()
        if any(text.startswith(prefix) for prefix in DEPRECATED_ASSESSMENT_NOTES):
            paragraph._element.getparent().remove(paragraph._element)


def format_run(run, *, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Times New Roman")


def set_paragraph_content(paragraph, value: Any, *, bold: bool = False) -> None:
    items = content_items(value) if not isinstance(value, (str, int)) else [str(value).strip()]
    clear_paragraph(paragraph)
    for index, item in enumerate(items):
        if index:
            paragraph.add_run().add_break()
        format_run(paragraph.add_run(item), bold=bold)


def set_cell_content(
    cell, value: Any, *, bullet_list: bool = False, render_emphasis: bool = False
) -> None:
    items = content_items(value) if not isinstance(value, (str, int)) else [str(value).strip()]
    cell.text = ""
    first = cell.paragraphs[0]
    for index, item in enumerate(items):
        paragraph = first if index == 0 else cell.add_paragraph()
        clear_paragraph(paragraph)
        if bullet_list:
            paragraph.style = "List Bullet"
        match = EMPHASIZED_ITEM.fullmatch(item) if render_emphasis else None
        if match:
            format_run(paragraph.add_run(match.group("text").strip()), bold=True, italic=True)
        else:
            format_run(paragraph.add_run(item))
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.insert(0, tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width_dxa))


def set_cell_margins(cell, value: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_repeating_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def mark_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:cantSplit"))
    if marker is None:
        marker = OxmlElement("w:cantSplit")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def remove_body_rows(table) -> None:
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)


def administrative_value(value: Any) -> str:
    return str(value).strip() if value is not None and str(value).strip() else "________________"


def fill_front_matter(doc, data: dict[str, Any]) -> None:
    language = data["language"]
    paragraph = doc.paragraphs[0]
    clear_paragraph(paragraph)
    format_run(paragraph.add_run(LABELS[language]["organization"]), bold=False)
    format_run(paragraph.add_run(administrative_value(data.get("organization"))), bold=False)

    table = doc.tables[0]
    values = {
        0: data["section"],
        1: administrative_value(data.get("teacher")),
        2: administrative_value(data.get("date")),
        3: data["grade"],
        5: data["topic"],
        6: data["learning_objectives"],
        7: data["lesson_objectives"],
        8: data["assessment_criteria"],
    }
    for row_index, value in values.items():
        set_cell_content(
            table.rows[row_index].cells[1],
            value,
            bullet_list=row_index in {7, 8},
        )
    set_cell_content(table.rows[4].cells[1], administrative_value(data.get("present")))
    set_cell_content(table.rows[4].cells[3], administrative_value(data.get("absent")))


def render_teacher_actions(cell, stage: dict[str, Any], language: str) -> None:
    cell.text = ""
    first = cell.paragraphs[0]
    entries = [
        (METHOD_PREFIX[language] + method["name"], True)
        for method in stage["methods"]
    ]
    bindings = stage.get("action_work_forms")
    if bindings is None:
        if len(stage.get("work_forms", [])) != 1:
            raise KSPError("Explicit action_work_forms required before rendering multiple work forms.")
        bindings = [stage["work_forms"] for _ in stage["teacher_actions"]]
    for action, forms in zip(stage["teacher_actions"], bindings):
        entries.extend((WORK_FORMS[language][form], True) for form in forms)
        entries.append((action, False))
    for index, (text, is_method) in enumerate(entries):
        paragraph = first if index == 0 else cell.add_paragraph()
        clear_paragraph(paragraph)
        format_run(paragraph.add_run(text), bold=is_method, italic=is_method and text.startswith(METHOD_PREFIX[language]))
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def rendered_task_columns(stage: dict[str, Any], language: str) -> tuple[list[str], list[str]]:
    learner_actions: list[str] = []
    assessment: list[str] = []
    for index, task in enumerate(stage["tasks"], 1):
        task_id = f"{TASK_LABEL[language]} {index}"
        learner_actions.append(f"{task_id}: {task['instruction']}")
        learner_actions.append(task['learner_action'])
        assessment.append(f"{task_id} — {scored_heading(language, task['total_points'])}")
        assessment.extend(f"• {item['text']} — {item['points']};" for item in task["descriptor_scores"])
        if task.get("support"):
            assessment.append(SUPPORT_LABEL[language] + task["support"])
        assessment.append(f"{LABELS[language]['feedback']}: {task['feedback']}")
    if not learner_actions:
        learner_actions = ["—"]
        assessment = ["—"]
    return learner_actions, assessment
def fill_stages(doc, data: dict[str, Any]) -> None:
    table = doc.tables[1]
    remove_body_rows(table)
    mark_repeating_header(table.rows[0])
    language = data["language"]
    for stage in data["stages"]:
        row = table.add_row()
        mark_row_cant_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, LESSON_WIDTHS_DXA[index])
            set_cell_margins(cell)
        set_cell_content(row.cells[0], f"{stage['name']} ({stage['minutes']} {LABELS[language]['minute']})")
        if data["schema_version"] == "3.0":
            render_teacher_actions(row.cells[1], stage, language)
            learner_items, assessment_items = rendered_task_columns(stage, language)
            set_cell_content(row.cells[2], learner_items)
            set_cell_content(row.cells[3], assessment_items)
            for paragraph in row.cells[3].paragraphs:
                text = paragraph.text
                if any(text == f"{TASK_LABEL[language]} {index} — {scored_heading(language, task['total_points'])}" for index, task in enumerate(stage["tasks"], 1)):
                    for run in paragraph.runs:
                        format_run(run, bold=True)
                elif text.startswith(SUPPORT_LABEL[language]):
                    clear_paragraph(paragraph)
                    format_run(paragraph.add_run(SUPPORT_LABEL[language]), bold=True)
                    format_run(paragraph.add_run(text[len(SUPPORT_LABEL[language]):]), bold=False)
        else:
            set_cell_content(row.cells[1], stage["teacher_actions"], render_emphasis=True)
            set_cell_content(row.cells[2], stage["learner_actions"])
            set_cell_content(row.cells[3], stage["assessment"])
        set_cell_content(row.cells[4], stage["resources"])


def add_heading(doc, text: str, level: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(3)
    format_run(paragraph.add_run(text), bold=True)


def add_appendix_content(doc, value: Any) -> None:
    items = content_items(value)
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet" if len(items) > 1 else None)
        paragraph.paragraph_format.space_after = Pt(3)
        format_run(paragraph.add_run(item))


def add_italic_notice(doc, heading: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = "Heading 2"
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    format_run(paragraph.add_run(heading), bold=True, italic=True)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    format_run(paragraph.add_run(text), italic=True)


def named_methods_in_content(value: Any) -> list[str]:
    """Return recognized professional method terms in one appendix content block."""
    text = "\n".join(content_items(value))
    return [
        name for name, pattern in METHOD_TERM_PATTERNS.items()
        if pattern.search(text)
    ]


def appendix_render_order(appendix: dict[str, Any]) -> list[str]:
    """Place term explanations immediately after the block that first uses a term."""
    order = [
        key for key in APPENDIX_ORDER
        if key in appendix and nonempty(appendix[key])
    ]
    if "term_explanations" not in appendix or not nonempty(appendix["term_explanations"]):
        return order

    anchor = next(
        (key for key in order if named_methods_in_content(appendix[key])),
        None,
    )
    if anchor is None:
        anchor = next(
            (key for key in ("methodology_application", "model_rationale", "goal_analysis") if key in order),
            order[-1] if order else None,
        )
    if anchor is None:
        return ["term_explanations"]
    order.insert(order.index(anchor) + 1, "term_explanations")
    return order


def append_methodological_appendix(doc, data: dict[str, Any]) -> None:
    remove_deprecated_template_notes(doc)
    # A standalone page-break paragraph can spill to a blank page after a full table.
    body = doc._element.body
    for node in list(body)[::-1]:
        if node.tag == qn("w:sectPr"):
            continue
        if node.tag == qn("w:p") and not node.xpath(".//w:t"):
            body.remove(node)
        else:
            break
    labels = LABELS[data["language"]]
    add_heading(doc, labels["appendix"], 1)
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    appendix = data["methodological_appendix"]
    for key in appendix_render_order(appendix):
        add_heading(doc, labels[key], 2)
        add_appendix_content(doc, appendix[key])
    add_italic_notice(doc, labels["professional_notice"], labels["professional_notice_text"])


def normalize_document(doc) -> None:
    for style in doc.styles:
        if hasattr(style, "font"):
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            if style.name.startswith("Heading"):
                style.font.bold = True
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            format_run(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        format_run(run)


def remove_kazakh_criteria_note(doc, language: str) -> None:
    if language != "kk":
        return
    removed_notes = {
        "Ескерту: «Бағалау критерийлері» — әдістемелік қосымша өріс.",
        "Ескерту: «Бағалау критерийлері» — әдістемелік қосымша бөлім.",
    }
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() in removed_notes:
            element = paragraph._element
            element.getparent().remove(element)


def normalize_ooxml(path: Path) -> None:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=path.parent) as handle:
        temp_path = Path(handle.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename in {"word/document.xml", "word/styles.xml"}:
                    root = etree.fromstring(payload)
                    for fonts in root.xpath(".//w:rFonts", namespaces=NS):
                        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                            fonts.set(f"{{{W_NS}}}{attr}", "Times New Roman")
                        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                            fonts.attrib.pop(f"{{{W_NS}}}{attr}", None)
                    for size in root.xpath(".//w:sz | .//w:szCs", namespaces=NS):
                        size.set(f"{{{W_NS}}}val", "24")
                    for height in root.xpath(".//w:trHeight[@w:hRule='exact']", namespaces=NS):
                        height.set(f"{{{W_NS}}}hRule", "atLeast")
                    payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                target.writestr(item, payload)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def all_text(doc) -> str:
    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def structural_audit(
    path: Path,
    expected_stages: int | None = None,
    expected_bullets: dict[int, int] | None = None,
    expected_stage_contract: list[dict[str, Any]] | None = None,
    expected_language: str | None = None,
) -> dict[str, Any]:
    if not path.exists() or path.stat().st_size == 0:
        raise KSPError(f"DOCX does not exist or is empty: {path}")
    doc = Document(path)
    failures: list[str] = []
    if len(doc.tables) < 2:
        failures.append("Expected at least two KSP tables.")
    else:
        if len(doc.tables[0].rows) < 9:
            failures.append("The KSP metadata table is incomplete.")
        else:
            for row_index, label in ((7, "lesson objectives"), (8, "assessment criteria")):
                paragraphs = [p for p in doc.tables[0].rows[row_index].cells[1].paragraphs if p.text.strip()]
                if not paragraphs or any(p.style.name != "List Bullet" for p in paragraphs):
                    failures.append(f"The {label} must be formatted as a bulleted list.")
                if any("\n" in p.text or "\r" in p.text for p in paragraphs):
                    failures.append(f"Each {label} item must be a separate bullet paragraph.")
                if any(re.match(r"^\s*(?:[•◦▪‣⁃*-]|\d+[.)])\s+", p.text) for p in paragraphs):
                    failures.append(f"The {label} contain typed markers instead of clean list items.")
                if expected_bullets is not None and len(paragraphs) != expected_bullets[row_index]:
                    failures.append(
                        f"The {label} bullet count is {len(paragraphs)}; "
                        f"expected {expected_bullets[row_index]}."
                    )
        if len(doc.tables[1].columns) != 5:
            failures.append("The lesson-progress table must have five columns.")
        if expected_stages is not None and len(doc.tables[1].rows) != expected_stages + 1:
            failures.append("The generated lesson-stage row count is incorrect.")
        for row in doc.tables[1].rows[1:]:
            if contains_framework(" ".join(cell.text for cell in row.cells)):
                failures.append("Move methodology frameworks from the lesson-progress table to the methodological appendix.")
        header_xml = doc.tables[1].rows[0]._tr.xml
        if "tblHeader" not in header_xml:
            failures.append("The lesson-progress table header is not marked to repeat.")
        if expected_stage_contract is not None and expected_language is not None:
            reference_doc = Document()
            reference_doc.add_table(rows=1, cols=2)
            reference_doc.add_table(rows=1, cols=5)
            fill_stages(reference_doc, {"schema_version": "3.0", "language": expected_language, "stages": expected_stage_contract})
            for actual, expected in zip(doc.tables[1].rows[1:], reference_doc.tables[1].rows[1:]):
                for column in (1, 3):
                    def signature(cell):
                        return [[(run.text, bool(run.bold), bool(run.italic)) for run in p.runs if run.text] for p in cell.paragraphs]
                    if signature(actual.cells[column]) != signature(expected.cells[column]):
                        failures.append("Teacher-action bindings or scored descriptor content/format differs from source.")
            prefix = METHOD_PREFIX[expected_language]
            for stage_index, (stage, row) in enumerate(
                zip(expected_stage_contract, doc.tables[1].rows[1:]), start=1
            ):
                paragraphs = [p for p in row.cells[1].paragraphs if p.text.strip()]
                expected_methods = [prefix + method["name"] for method in stage["methods"]]
                expected_actions = content_items(stage["teacher_actions"])
                method_paragraphs = [p for p in paragraphs if p.text.startswith(prefix)]
                if [p.text for p in method_paragraphs] != expected_methods:
                    failures.append(
                        f"Stage {stage_index} {stage['name']!r}: method lines are missing, reordered, or use the wrong localized prefix."
                    )
                for paragraph in method_paragraphs:
                    visible_runs = [run for run in paragraph.runs if run.text]
                    if not visible_runs or any(run.bold is not True or run.italic is not True for run in visible_runs):
                        failures.append(
                            f"Stage {stage_index} {stage['name']!r}: method name is not fully bold italic."
                        )
                action_paragraphs = [p for p in paragraphs if p.text in expected_actions]
                if [p.text for p in action_paragraphs] != expected_actions:
                    failures.append(
                        f"Stage {stage_index} {stage['name']!r}: teacher actions are missing or not in separate paragraphs."
                    )
                for paragraph in action_paragraphs:
                    if any(run.bold is True or run.italic is True for run in paragraph.runs if run.text):
                        failures.append(
                            f"Stage {stage_index} {stage['name']!r}: ordinary teacher action must use regular type."
                        )
        else:
            teacher_method_runs = [
                run
                for row in doc.tables[1].rows[1:]
                for paragraph in row.cells[1].paragraphs
                for run in paragraph.runs
                if run.text.strip() and run.bold is True and run.italic is True
            ]
            if not teacher_method_runs:
                failures.append("Legacy document has no bold-italic method or technique label.")

    raw_text = all_text(doc)
    if "***" in raw_text:
        failures.append("Markdown emphasis markers leaked into the DOCX.")
    text = raw_text.casefold()
    label_text = "\n".join(
        row.cells[0].text.casefold()
        for row in doc.tables[0].rows
        if row.cells
    ) if doc.tables else ""
    found = [term for term in BANNED if term in label_text]
    if found:
        failures.append("Banned terminology found: " + ", ".join(found))
    if not any(label["appendix"].casefold() in text for label in LABELS.values()):
        failures.append("The methodological appendix is missing.")
    for key in ("knowledge_skills_analysis", "methodology_application"):
        if not any(label[key].casefold() in text for label in LABELS.values()):
            failures.append(f"The methodological appendix section is missing: {key}.")
    if not any(label["professional_notice_text"].casefold() in text for label in LABELS.values()):
        failures.append("The fixed professional notice is missing.")
    if any(old in text for old in ("дереккөздер және мәртебесі", "источники и статус", "sources and status")):
        failures.append("The sources heading must be exactly Дереккөздер / Источники / Sources.")
    notice_texts = {
        value.casefold()
        for label in LABELS.values()
        for value in (label["professional_notice"], label["professional_notice_text"])
    }
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().casefold() in notice_texts:
            visible_runs = [run for run in paragraph.runs if run.text.strip()]
            if not visible_runs or any(run.italic is not True for run in visible_runs):
                failures.append("The professional notice heading and text must be italic.")
    if any(old in text for old in ("ішкі аудит қорытындысы", "итог внутреннего аудита", "internal-audit summary")):
        failures.append("The internal-audit summary must not appear in the teacher-facing DOCX.")

    with zipfile.ZipFile(path, "r") as archive:
        for member in ("word/document.xml", "word/styles.xml"):
            root = etree.fromstring(archive.read(member))
            for fonts in root.xpath(".//w:rFonts", namespaces=NS):
                for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                    value = fonts.get(f"{{{W_NS}}}{attr}")
                    if value and value != "Times New Roman":
                        failures.append(f"Non-Times New Roman font in {member}: {value}")
                        break
            for size in root.xpath(".//w:sz | .//w:szCs", namespaces=NS):
                if size.get(f"{{{W_NS}}}val") != "24":
                    failures.append(f"Non-12 pt text size in {member}.")
                    break

    if failures:
        raise KSPError("Structural audit failed:\n- " + "\n- ".join(dict.fromkeys(failures)))
    return {
        "path": str(path),
        "pages": "visual render required",
        "tables": len(doc.tables),
        "lesson_stage_rows": len(doc.tables[1].rows) - 1,
        "font": "Times New Roman 12 pt",
        "status": "structural audit passed",
    }


def build(data: dict[str, Any], output: Path) -> dict[str, Any]:
    if data.get("schema_version") != "3.0":
        raise KSPError("Legacy input cannot be rendered without explicit scored descriptors and work-form migration to schema 3.0.")
    data = validate_input(data)
    template = TEMPLATES[data["language"]]
    if not template.exists():
        raise KSPError(f"Language template not found: {template}")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(template)
    if len(doc.tables) < 2 or len(doc.tables[0].rows) < 9:
        raise KSPError("The selected template has an unexpected structure.")
    fill_front_matter(doc, data)
    fill_stages(doc, data)
    remove_kazakh_criteria_note(doc, data["language"])
    append_methodological_appendix(doc, data)
    normalize_document(doc)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = (
        f"language={data['language']};lesson-duration={data['lesson_duration_minutes']};"
        f"requested-lesson-count={data['lesson_count']};"
        f"stage-total={sum(stage['minutes'] for stage in data['stages'])};"
        f"schema={data['schema_version']};validation={data['validation_status']}"
    )
    doc.save(output)
    normalize_ooxml(output)
    return structural_audit(
        output,
        expected_stages=len(data["stages"]),
        expected_bullets={
            7: len(data["lesson_objectives"]),
            8: len(data["assessment_criteria"]),
        },
        expected_stage_contract=data["stages"] if data["schema_version"] == "3.0" else None,
        expected_language=data["language"] if data["schema_version"] == "3.0" else None,
    )


def example(*, test_fixture: bool = False) -> dict[str, Any]:
    data = {
        "schema_version": "3.0",
        "intake_verification": {
            "confirmed_by_user": test_fixture,
            "confirmed_fields": sorted(REQUIRED_INTAKE_FIELDS),
            "source_summary": (
                "Synthetic developer test fixture."
                if test_fixture else
                "Replace with real teacher messages or attached-document evidence."
            ),
        },
        "instruction_language": "ru",
        "language": "ru",
        "subject": "Алгебра",
        "organization": "",
        "teacher": "",
        "date": "",
        "grade": "7",
        "class_size": 24,
        "teacher_experience": "8 лет",
        "qualification_category": "педагог-эксперт",
        "class_characteristics": "Разный темп работы; части учащихся нужна визуальная опора.",
        "present": "",
        "absent": "",
        "section": "Функция и график функции",
        "topic": "Линейная функция и её график",
        "lesson_count": 1,
        "learning_objectives": [
            "7.4.1.4 — знать определение линейной функции, строить её график и определять расположение в зависимости от коэффициента k"
        ],
        "lesson_objectives": [
            "Строить график линейной функции и объяснять влияние коэффициента k на его расположение."
        ],
        "assessment_criteria": [
            "Строит график по заданной формуле.",
            "Объясняет расположение графика с опорой на значение коэффициента k."
        ],
        "stages": [
            {
                "name": "Начало урока", "minutes": 5, "activity_type": "learning",
                "methods": [{"method_id": "method-01", "type": "supporting", "name": "Вспомни без подсказки"}],
                "teacher_actions": ["Организует актуализацию необходимых знаний."],
                "tasks": [{
                    "canonical_task_id": "task-01", "task_type": "oral",
                    "objective_refs": ["7.4.1.4"],
                    "instruction": "Объясните, как построить график линейной функции по двум точкам.",
                    "learner_action": "Воспроизводит алгоритм построения графика без подсказки.",
                    "expected_product": "Устное объяснение последовательности построения.",
                    "descriptor": "Называет не менее двух корректных шагов построения графика.",
                    "feedback": "Учитель уточняет пропущенный шаг вопросом и предлагает исправить ответ."
                }],
                "resources": "Доска."
            },
            {
                "name": "Основная часть", "minutes": 30, "activity_type": "learning",
                "methods": [{"method_id": "method-02", "type": "primary", "name": "Предскажи — проверь — объясни"}],
                "teacher_actions": ["Организует исследование и практику построения графиков."],
                "tasks": [{
                    "canonical_task_id": "task-02", "task_type": "written",
                    "objective_refs": ["7.4.1.4"],
                    "instruction": "Постройте два графика с разными значениями коэффициента k, сравните их и объясните различие.",
                    "learner_action": "Строит, сравнивает и объясняет графики.",
                    "expected_product": "Два графика и письменный вывод о влиянии коэффициента k.",
                    "descriptor": "Строит оба графика без ошибок и связывает их расположение со значением k.",
                    "feedback": "Учитель указывает, какой элемент графика или объяснения нужно перепроверить."
                }],
                "resources": "Карточки, координатная плоскость."
            },
            {
                "name": "Рефлексия", "minutes": 5, "activity_type": "learning",
                "methods": [{"method_id": "method-03", "type": "technique", "name": "Рефлексия по критерию"}],
                "teacher_actions": ["Организует итоговую рефлексию с доказательством достижения цели."],
                "tasks": [{
                    "canonical_task_id": "task-03", "task_type": "reflection",
                    "objective_refs": ["7.4.1.4"],
                    "instruction": "Сформулируйте вывод о достижении цели и назовите следующий шаг.",
                    "learner_action": "Соотносит свой результат с критерием и определяет следующий шаг.",
                    "expected_product": "Индивидуальный рефлексивный ответ.",
                    "descriptor": "Приводит одно доказательство достижения цели и формулирует конкретный следующий шаг.",
                    "feedback": "Учитель подтверждает доказательство или просит конкретизировать следующий шаг."
                }],
                "resources": "Карточка рефлексии."
            }
        ],
        "methodological_appendix": {
            "goal_analysis": "Цель требует построения графика и объяснения влияния коэффициента.",
            "knowledge_skills_analysis": [
                "ПРЕДМЕТНОЕ СОДЕРЖАНИЕ: линейная функция, её график и коэффициент k — основание: цель обучения и тема.",
                "ФАКТИЧЕСКИЕ ЗНАНИЯ: обозначение коэффициента k и координаты точек графика — основание: цель обучения.",
                "ПРЕДМЕТНЫЕ НАВЫКИ: строить и сопоставлять графики линейных функций — основание: цель обучения и задание урока.",
                "ПОЗНАВАТЕЛЬНЫЕ НАВЫКИ: применять способ построения, анализировать зависимость и объяснять вывод — основание: действия ученика."
            ],
            "methodology_application": "Управляемое исследование: ученики строят и сопоставляют графики, учитель собирает объяснения и при ошибке возвращает опору на координатную сетку и контрастный пример.",
            "model_rationale": "Выбрано управляемое исследование с последующей самостоятельной практикой.",
            "differentiation": "Опорная сетка доступна всем; подсказка выдаётся по диагностике; усложнение требует обоснования.",
            "formative_assessment": "Учитель собирает графики и объяснения, затем корректирует практику.",
            "alternatives": "Резерв: сопоставление формул и готовых графиков.",
            "audit_summary": "Цели, критерии, задания и время согласованы."
        }
    }
    for stage in data["stages"]:
        stage["work_forms"] = ["individual"]
        for task in stage["tasks"]:
            task["descriptor_scores"] = [{"text": task["descriptor"], "points": 1}]
    if not test_fixture:
        data["intake_verification"]["confirmed_fields"] = []
    return data


def parse_args(argv: Iterable[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", nargs="?", type=Path, help="UTF-8 JSON input")
    parser.add_argument("output", nargs="?", type=Path, help="output DOCX path")
    parser.add_argument("--check", type=Path, help="run structural audit on an existing DOCX")
    parser.add_argument(
        "--print-example", action="store_true",
        help="print a schema fixture that still requires real teacher confirmation",
    )
    return parser.parse_args(argv)


def main(argv: Iterable[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.print_example:
            print(json.dumps(example(), ensure_ascii=False, indent=2))
            return 0
        if args.check:
            print(json.dumps(structural_audit(args.check), ensure_ascii=False, indent=2))
            return 0
        if not args.input or not args.output:
            raise KSPError("Provide input.json and output.docx, or use --print-example/--check.")
        with args.input.open("r", encoding="utf-8") as handle:
            data = validate_input(json.load(handle))
        result = build(data, args.output.resolve())
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (KSPError, json.JSONDecodeError, OSError, zipfile.BadZipFile) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
