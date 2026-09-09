"""Build a scored worksheet from a strict lesson handoff (grades 1–11)."""

import argparse
import importlib.util
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Mm, Pt

from validate_derivative_alignment import validate_alignment

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("worksheet_ksp_layout", ROOT / "skills/school-lesson-planner/secondary/scripts/build_ksp.py")
KSP = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(KSP)
LABELS = {
    "ru": ("Рабочий лист", "Задание", "Руководство для учителя", "Страница {n} из {total}"),
    "kk": ("Жұмыс парағы", "Тапсырма", "Мұғалімге нұсқаулық", "{n}-бет / {total}"),
    "en": ("Worksheet", "Task", "Guide for the Teacher", "Page {n} of {total}"),
}


def build(handoff, manifest, language, output):
    validate_alignment(handoff, manifest)
    if manifest.get("artifact_type") != "worksheet" or language not in LABELS:
        raise ValueError("A worksheet manifest and kk/ru/en language are required")
    if handoff.get("language") != language:
        raise ValueError("Worksheet language must match the source lesson")
    labels = LABELS[language]
    doc = Document()
    tasks = {t["canonical_task_id"]: t for t in handoff["canonical_tasks"]}
    for index, item in enumerate(manifest["items"], 1):
        section = doc.sections[0] if index == 1 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(18)
        section.left_margin = section.right_margin = Mm(20)
        section.footer.is_linked_to_previous = False
        section.footer.paragraphs[0].text = labels[3].format(n=index, total=len(manifest["items"]))
        task = tasks[item["canonical_task_id"]]
        doc.add_paragraph(labels[0]).runs[0].bold = True
        doc.add_paragraph(f"{labels[1]} {index}").runs[0].bold = True
        doc.add_paragraph(task["instruction"])
        for _ in range(3):
            doc.add_paragraph("________________________________________________________________")
        doc.add_paragraph(KSP.scored_heading(language, item["total_points"])).runs[0].bold = True
        for score in item["descriptor_scores"]:
            doc.add_paragraph(f"• {score['text']} — {score['points']};")
        if item.get("support"):
            p = doc.add_paragraph()
            p.add_run(KSP.SUPPORT_LABEL[language]).bold = True
            p.add_run(item["support"])
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = ""
    doc.add_paragraph(labels[2]).runs[0].bold = True
    for index, item in enumerate(manifest["items"], 1):
        task = tasks[item["canonical_task_id"]]
        doc.add_paragraph(f"{labels[1]} {index}").runs[0].bold = True
        doc.add_paragraph(task["expected_product"])
        doc.add_paragraph(task["feedback"])
    KSP.normalize_document(doc)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    KSP.normalize_ooxml(output)
    audit(output, handoff, manifest, language)


def audit(path, handoff, manifest, language):
    validate_alignment(handoff, manifest)
    doc = Document(path)
    paragraphs = doc.paragraphs
    starts = [i for i, p in enumerate(paragraphs) if p.text == LABELS[language][0]]
    if len(starts) != len(manifest["items"]):
        raise ValueError("Worksheet task count differs from source")
    for index, item in enumerate(manifest["items"]):
        end = starts[index + 1] if index + 1 < len(starts) else next(i for i, p in enumerate(paragraphs) if p.text == LABELS[language][2])
        block = paragraphs[starts[index]:end]
        expected = [KSP.scored_heading(language, item["total_points"])] + [f"• {s['text']} — {s['points']};" for s in item["descriptor_scores"]]
        for text in expected:
            if sum(p.text == text for p in block) != 1:
                raise ValueError("Missing or changed scored descriptor in DOCX")
        heading = next(p for p in block if p.text == expected[0])
        if not all(r.bold for r in heading.runs if r.text):
            raise ValueError("Descriptor heading must be bold")
        if item.get("support"):
            support = next((p for p in block if p.text == KSP.SUPPORT_LABEL[language] + item["support"]), None)
            if support is None or not support.runs[0].bold:
                raise ValueError("Support text/format differs from source")
    return {"status": "PASS", "tasks": len(starts)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8"))
    build(data["lesson_plan_handoff"], data["derivative"], data["language"], args.output)
