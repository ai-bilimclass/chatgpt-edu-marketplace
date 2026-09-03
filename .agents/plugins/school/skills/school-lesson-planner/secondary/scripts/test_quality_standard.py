from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "references" / "lesson-quality-standard.docx"


def all_text():
    doc = Document(DOCX)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def test_dual_forms_and_scales():
    text = all_text()
    assert "Форма 1. Качество проектирования КСП" in text
    assert "Форма 2. Качество проведённого урока" in text
    assert "Максимум: 28 баллов" in text
    assert "Максимум: 36 баллов" in text


def test_required_explicit_criteria_and_adaptation_wording():
    text = all_text()
    assert "Связность «цель → задание → результат → оценивание»" in text
    assert "Завершение с демонстрацией результата" in text
    assert "содержательная цель не снижается без нормативно и педагогически обоснованной необходимости" in text


def test_three_levels_gates_and_source_status():
    text = all_text()
    assert "Уровень 3. Решение для следующего урока" in text
    for gate in ["Безопасность", "Недискриминация", "Предметная корректность", "Защита данных", "Ответственное использование ИИ"]:
        assert gate in text
    assert "Печатные страницы 104–118 и 174" in text
    assert "Не является нормативным правовым актом" in text
