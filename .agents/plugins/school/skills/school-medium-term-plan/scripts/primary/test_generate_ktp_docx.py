#!/usr/bin/env python3
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

from generate_ktp_docx import (
    EXPECTED_QUARTER_DATES,
    EXPECTED_WEEKS,
    audit_docx_structure,
    build,
    validate,
)

PRIMARY_SUMMATIVE_MARKERS = ("БЖБ", "ТЖБ", "СОР", "СОЧ", "SAU", "SAT")


def assessment_values(data: dict) -> list[str]:
    return [
        str(lesson.get("assessment_type", ""))
        for quarter in data["quarters"]
        for lesson in quarter["lessons"]
    ]


def contains_primary_summative_marker(value: str) -> bool:
    return any(marker in value.upper() for marker in PRIMARY_SUMMATIVE_MARKERS)


def make_data(hours_per_week: int) -> dict:
    weekdays = [1 + (index % 4) for index in range(hours_per_week)]
    quarters = []
    number = 1
    for index, weeks in enumerate(EXPECTED_WEEKS):
        count = weeks * hours_per_week
        lessons = []
        for position in range(count):
            lessons.append({
                "number": number,
                "section": f"Раздел {index + 1}",
                "topic": "Официальное содержание",
                "learning_objectives": "2.1.1.1 – описывать официальное содержание учебной программы",
                "hours": 1,
                "assessment_type": "",
                "note": "",
            })
            number += 1
        quarters.append({
            "name": f"{index + 1} четверть",
            "weeks": weeks,
            "soch_required": True,
            "start_date": EXPECTED_QUARTER_DATES[index][0],
            "end_date": EXPECTED_QUARTER_DATES[index][1],
            "lessons": lessons,
        })
    return {
        "subject": "Математика",
        "grade": 2,
        "education_level": "primary",
        "instruction_language": "ru",
        "profile_direction": "not_applicable",
        "curriculum_order": "№399",
        "curriculum_appendix": "52",
        "curriculum_revision_date": "2026-08-14",
        "curriculum_plan_order": "№500",
        "curriculum_plan_variant": "Тестовый вариант типового учебного плана",
        "curriculum_plan_revision_date": "2026-08-14",
        "subject_profile": "mathematics_informatics",
        "language": "ru",
        "requested_language": "ru",
        "objectives_language": "ru",
        "language_mismatch_confirmed": False,
        "academic_year": "2026–2027",
        "total_hours": 34 * hours_per_week,
        "hours_per_week": hours_per_week,
        "lesson_weekdays": weekdays,
        "extended_objectives": [],
        "content_source": "Официальная тестовая программа",
        "content_source_type": "teacher_uploaded_program",
        "teacher_program_file": "teacher-curriculum.pdf",
        "cross_cutting_themes": ["Раздел 1", "Раздел 2", "Раздел 3", "Раздел 4"],
        "lesson_topics_confirmed_by_teacher": True,
        "official_program_verified": True,
        "source_content_complete": True,
        "source_conflicts": [],
        "calendar_verified": True,
        "calendar_source": "Официальный тестовый календарь",
        "non_instruction_dates": [],
        "public_holidays": [],
        "summative_assessment_required": True,
        "quarters": quarters,
    }


class DynamicHoursValidationTest(unittest.TestCase):
    def test_accepts_zero_width_formatting_inside_objective_code(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["learning_objectives"] = (
            "1.\u200b1.\u200b1.\u200b1 есту қабілетіне сүйеніп таныс дыбыстарды тану "
            "және мәтін\u200bаралық белгі"
        )

        validate(data)

        normalized = data["quarters"][0]["lessons"][0]["learning_objectives"]
        self.assertTrue(normalized.startswith("1.1.1.1 "))
        self.assertIn("мәтін\u200bаралық", normalized)

    def test_accepts_complete_teacher_objectives_without_uploaded_program(self):
        data = make_data(1)
        data["content_source_type"] = "teacher_provided_objectives"
        data["teacher_program_file"] = ""
        data["official_program_verified"] = False
        validate(data)

    def test_rejects_primary_plan_without_cross_cutting_themes(self):
        data = make_data(1)
        data["cross_cutting_themes"] = []
        with self.assertRaisesRegex(ValueError, "content organizers"):
            validate(data)

    def test_accepts_grade_one_bukvar_with_teacher_sections_and_topics(self):
        data = make_data(1)
        data["grade"] = 1
        data["subject"] = "Букварь"
        data["subject_profile"] = "languages_literature"
        data["total_hours"] = 33
        data["quarters"][2]["weeks"] = 9
        data["quarters"][2]["lessons"].pop()
        data["content_source_type"] = "teacher_provided_sections_and_topics"
        data["teacher_program_file"] = "teacher-pasted-sections-topics.txt"
        data["official_program_verified"] = False
        data["cross_cutting_themes"] = []
        data["content_organizer_type"] = "sections"
        data["content_organizers"] = ["Раздел 1", "Раздел 2", "Раздел 3", "Раздел 4"]
        number = 1
        for quarter in data["quarters"]:
            for lesson in quarter["lessons"]:
                lesson["number"] = number
                lesson["learning_objectives"] = "1.1.1.1 – точная цель из источника учителя"
                number += 1
        validate(data)

    def test_rejects_grade_one_alippe_with_cross_cutting_themes(self):
        data = make_data(1)
        data["grade"] = 1
        data["subject"] = "Әліппе"
        data["subject_profile"] = "languages_literature"
        data["total_hours"] = 33
        data["quarters"][2]["weeks"] = 9
        data["quarters"][2]["lessons"].pop()
        with self.assertRaisesRegex(ValueError, "teacher-provided sections"):
            validate(data)

    def test_rejects_unconfirmed_primary_lesson_topics(self):
        data = make_data(1)
        data["lesson_topics_confirmed_by_teacher"] = False
        with self.assertRaisesRegex(ValueError, "explicitly confirmed or corrected"):
            validate(data)

    def test_rejects_lesson_outside_confirmed_cross_cutting_themes(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["section"] = "Белгісіз ортақ тақырып"
        with self.assertRaisesRegex(ValueError, "not listed in content organizers"):
            validate(data)

    def test_rejects_unsupported_primary_source(self):
        data = make_data(1)
        data["content_source_type"] = "model_memory"
        data["teacher_program_file"] = ""
        data["official_program_verified"] = False
        with self.assertRaisesRegex(ValueError, "teacher-uploaded program or a complete set"):
            validate(data)

    def test_official_calendar_reference_contains_order_and_holidays(self):
        root = Path(__file__).parents[2]
        calendar = (root / "references" / "primary" / "academic-calendar-2026-2027.md").read_text(encoding="utf-8")
        holidays = json.loads((root / "references" / "primary" / "official-holidays-2026-2027.json").read_text(encoding="utf-8"))
        self.assertIn("№213", calendar)
        self.assertIn("08.02.2027–14.02.2027", calendar)
        dates = {item["date"] for item in holidays["public_holidays"]}
        self.assertTrue({"2026-12-16", "2027-03-15", "2027-05-10"}.issubset(dates))
        self.assertEqual(holidays["pending_variable_holidays"][0]["status"], "not_officially_published_as_of_verification_date")

    def test_accepts_grade_one_33_week_calendar_without_assessments(self):
        data = make_data(1)
        data["grade"] = 1
        data["total_hours"] = 33
        data["summative_assessment_required"] = False
        data["quarters"][2]["weeks"] = 9
        data["quarters"][2]["lessons"].pop()
        number = 1
        for quarter in data["quarters"]:
            for lesson in quarter["lessons"]:
                lesson["number"] = number
                lesson["learning_objectives"] = "1.1.1.1 – ресми оқу мақсаты"
                number += 1
        validate(data)
        self.assertEqual(sum(len(q["lessons"]) for q in data["quarters"]), 33)

    def test_clears_assessments_when_summative_not_required(self):
        data = make_data(1)
        data["summative_assessment_required"] = False
        data["quarters"][0]["lessons"][0]["assessment_type"] = "СОР №9"
        validate(data)
        self.assertTrue(all(
            lesson["assessment_type"] == ""
            for quarter in data["quarters"] for lesson in quarter["lessons"]
        ))
        self.assertTrue(all(not quarter["soch_required"] for quarter in data["quarters"]))

    def test_primary_mode_clears_assessments_even_when_legacy_input_requests_them(self):
        data = make_data(1)
        data["summative_assessment_required"] = True
        data["quarters"][0]["lessons"][0]["assessment_type"] = "БЖБ №1"
        data["quarters"][0]["lessons"][1]["assessment_type"] = "ТЖБ"
        validate(data)
        self.assertFalse(data["summative_assessment_required"])
        self.assertTrue(all(not quarter["soch_required"] for quarter in data["quarters"]))
        self.assertTrue(all(
            lesson["assessment_type"] == ""
            for quarter in data["quarters"] for lesson in quarter["lessons"]
        ))

    def test_primary_docx_never_displays_secondary_summative_labels(self):
        data = make_data(1)
        data["summative_assessment_required"] = True
        data["quarters"][0]["lessons"][0]["assessment_type"] = "СОР №1"
        data["quarters"][0]["lessons"][1]["assessment_type"] = "СОЧ"
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "primary-no-summative-labels.docx"
            build(data, output)
            document = Document(output)
            text = "\n".join(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
        for forbidden in ("БЖБ", "ТЖБ", "СОР", "СОЧ", "SAU", "SAT"):
            self.assertNotIn(forbidden, text)

    def test_trilingual_primary_document_does_not_require_assessment_language(self):
        data = make_data(1)
        data.update({
            "language": "trilingual",
            "requested_language": "trilingual",
            "objectives_language": "mixed",
        })
        data.pop("assessment_language", None)
        validate(data)
        self.assertFalse(data["summative_assessment_required"])

    def test_requires_order_500_context(self):
        data = make_data(1)
        data["curriculum_plan_order"] = "№399"
        with self.assertRaisesRegex(ValueError, "№500"):
            validate(data)

    def test_trilingual_document_clears_selected_assessment_language_labels(self):
        data = make_data(1)
        data.update({
            "language": "trilingual",
            "requested_language": "trilingual",
            "objectives_language": "mixed",
            "instruction_language": "mixed",
            "assessment_language": "en",
        })
        validate(data)
        self.assertFalse(any(assessment_values(data)))
        self.assertFalse(data["summative_assessment_required"])

    def test_trilingual_document_does_not_require_assessment_language(self):
        data = make_data(1)
        data.update({
            "language": "trilingual",
            "requested_language": "trilingual",
            "objectives_language": "mixed",
            "instruction_language": "mixed",
        })
        data.pop("assessment_language", None)
        validate(data)
        self.assertFalse(any(assessment_values(data)))

    def test_kazakh_primary_ktp_contains_no_summative_labels(self):
        data = make_data(1)
        data.update({
            "language": "kk",
            "requested_language": "kk",
            "objectives_language": "kk",
            "instruction_language": "kk",
        })
        validate(data)
        self.assertFalse(any(assessment_values(data)))

    def test_english_primary_ktp_contains_no_summative_labels(self):
        data = make_data(1)
        data.update({
            "language": "en",
            "requested_language": "en",
            "objectives_language": "en",
            "instruction_language": "en",
        })
        for quarter_index, quarter in enumerate(data["quarters"], start=1):
            quarter["name"] = f"Term {quarter_index}"
            for lesson in quarter["lessons"]:
                lesson["section"] = f"Section {quarter_index}"
                lesson["topic"] = "Official curriculum content"
                lesson["learning_objectives"] = (
                    "2.1.1.1 – describe official curriculum content"
                )
        data["cross_cutting_themes"] = [f"Section {index}" for index in range(1, 5)]
        validate(data)
        self.assertFalse(any(assessment_values(data)))

    def test_clears_mixed_manual_markers_in_english_primary_ktp(self):
        data = make_data(1)
        data.update({
            "language": "en",
            "requested_language": "en",
            "objectives_language": "en",
            "instruction_language": "en",
        })
        data["quarters"][0]["lessons"][0]["assessment_type"] = "СОР №9"
        data["quarters"][0]["lessons"][1]["assessment_type"] = "ТЖБ"
        validate(data)
        self.assertFalse(any(assessment_values(data)))

    def test_builds_english_docx_with_headers_and_without_summative_markers(self):
        data = make_data(1)
        data.update({
            "subject": "Mathematics",
            "language": "en",
            "requested_language": "en",
            "objectives_language": "en",
            "instruction_language": "en",
            "content_source": "Official curriculum",
            "calendar_source": "Official academic calendar",
        })
        for quarter_index, quarter in enumerate(data["quarters"], start=1):
            quarter["name"] = f"Term {quarter_index}"
            for lesson in quarter["lessons"]:
                lesson["section"] = f"Section {quarter_index}"
                lesson["topic"] = "Official curriculum content"
                lesson["learning_objectives"] = (
                    "2.1.1.1 – describe official curriculum content"
                )
        data["cross_cutting_themes"] = [f"Section {index}" for index in range(1, 5)]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "english-ktp.docx"
            build(data, output)
            document = Document(output)
            self.assertIn("Medium-term (calendar-thematic) subject plan", document.paragraphs[0].text)
            self.assertEqual(
                [cell.text.strip() for cell in document.tables[0].rows[0].cells],
                ["No.", "Section /\nCross-cutting themes", "Lesson topic", "Learning objectives",
                 "Number of hours", "Dates", "Notes"],
            )
            notes = [row.cells[6].text.strip() for row in document.tables[0].rows]
            self.assertFalse(any(contains_primary_summative_marker(note) for note in notes))

    def test_requires_explicit_choice_when_english_and_objective_languages_differ(self):
        data = make_data(1)
        data.update({
            "language": "en",
            "requested_language": "en",
            "objectives_language": "kk",
            "instruction_language": "en",
        })
        with self.assertRaisesRegex(ValueError, "ағылшынша"):
            validate(data)

    def test_accepts_english_after_explicit_mismatch_choice_without_assessment_labels(self):
        data = make_data(1)
        data.update({
            "language": "en",
            "requested_language": "ru",
            "objectives_language": "en",
            "instruction_language": "en",
            "language_mismatch_confirmed": True,
        })
        validate(data)
        self.assertFalse(any(assessment_values(data)))

    def test_rejects_subject_profile_mismatch(self):
        data = make_data(1)
        data["subject_profile"] = "natural_sciences"
        with self.assertRaisesRegex(ValueError, "subject_profile must be"):
            validate(data)

    def test_rejects_non_primary_grade(self):
        data = make_data(1)
        data["grade"] = 5
        with self.assertRaisesRegex(ValueError, "1 to 4"):
            validate(data)

    def test_requires_primary_education_level(self):
        data = make_data(1)
        data["education_level"] = "basic_secondary"
        with self.assertRaisesRegex(ValueError, "education_level must be primary"):
            validate(data)

    def test_requires_safety_metadata_for_laboratory_lesson(self):
        data = make_data(1)
        data["subject"] = "Физика"
        data["subject_profile"] = "natural_sciences"
        lesson = data["quarters"][0]["lessons"][0]
        lesson["lesson_mode"] = "laboratory"
        lesson["practical_work"] = "Тығыздықты анықтау"
        with self.assertRaisesRegex(ValueError, "safety_required"):
            validate(data)

    def test_accepts_complete_laboratory_metadata(self):
        data = make_data(1)
        data["subject"] = "Физика"
        data["subject_profile"] = "natural_sciences"
        lesson = data["quarters"][0]["lessons"][0]
        lesson.update({
            "lesson_mode": "laboratory",
            "practical_work": "Тығыздықты анықтау",
            "safety_required": True,
        })
        validate(data)

    def test_accepts_dynamic_hours_without_upper_limit(self):
        for hours_per_week in (1, 2, 10, 11, 15):
            with self.subTest(hours_per_week=hours_per_week):
                validate(make_data(hours_per_week))

    def test_rejects_fixed_68_hours_for_one_hour_subject(self):
        data = make_data(1)
        data["total_hours"] = 68
        with self.assertRaisesRegex(ValueError, "do not equal total hours"):
            validate(data)

    def test_clears_manual_sor_in_primary_ktp(self):
        data = make_data(2)
        data["quarters"][0]["lessons"][0]["assessment_type"] = "СОР"
        validate(data)
        self.assertFalse(any(assessment_values(data)))

    def test_section_boundaries_do_not_create_assessment_labels(self):
        data = make_data(1)
        lessons = data["quarters"][0]["lessons"]
        for index, lesson in enumerate(lessons):
            lesson["section"] = "Бөлім A" if index < 3 else "Бөлім B"
        data["cross_cutting_themes"] = [
            "Бөлім A", "Бөлім B", "Раздел 2", "Раздел 3", "Раздел 4"
        ]
        validate(data)
        self.assertFalse(any(lesson["assessment_type"] for lesson in lessons))

    def test_docx_does_not_display_section_assessment_labels(self):
        data = make_data(1)
        lessons = data["quarters"][0]["lessons"]
        for index, lesson in enumerate(lessons):
            lesson["section"] = "Бөлім A" if index < 3 else "Бөлім B"
        data["cross_cutting_themes"] = [
            "Бөлім A", "Бөлім B", "Раздел 2", "Раздел 3", "Раздел 4"
        ]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "section-assessments.docx"
            build(data, output)
            document = Document(output)
            notes = [row.cells[6].text.strip() for row in document.tables[0].rows]
            self.assertFalse(any(contains_primary_summative_marker(note) for note in notes))

    def test_all_quarters_remain_free_of_assessment_labels(self):
        data = make_data(1)
        validate(data)
        for quarter in data["quarters"]:
            self.assertFalse(any(
                lesson["assessment_type"] for lesson in quarter["lessons"]
            ))

    def test_continuing_section_across_quarters_creates_no_assessment_marker(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][-1]["section"] = "Жалғасатын бөлім"
        data["quarters"][1]["lessons"][0]["section"] = "Жалғасатын бөлім"
        data["cross_cutting_themes"].append("Жалғасатын бөлім")
        validate(data)
        self.assertEqual(data["quarters"][0]["lessons"][-1]["assessment_type"], "")
        self.assertEqual(data["quarters"][1]["lessons"][0]["assessment_type"], "")

    def test_cross_quarter_section_does_not_create_sor_or_soch(self):
        data = make_data(1)
        lessons = data["quarters"][0]["lessons"]
        lessons[-1]["section"] = "Келесі бөлім"
        data["quarters"][1]["lessons"][0]["section"] = "Келесі бөлім"
        data["cross_cutting_themes"].append("Келесі бөлім")
        validate(data)
        self.assertFalse(any(assessment_values(data)))

    def test_does_not_insert_soch_and_preserves_lesson_content(self):
        data = make_data(2)
        lessons = data["quarters"][0]["lessons"]
        lesson_count = len(lessons)
        quarter_hours = sum(lesson["hours"] for lesson in lessons)
        penultimate_topic = lessons[-2]["topic"]
        penultimate_objective = lessons[-2]["learning_objectives"]
        validate(data)
        self.assertEqual(lessons[-2]["assessment_type"], "")
        self.assertEqual(lessons[-2]["topic"], penultimate_topic)
        self.assertEqual(lessons[-2]["learning_objectives"], penultimate_objective)
        self.assertEqual(lessons[-2]["hours"], 1)
        self.assertEqual(len(lessons), lesson_count)
        self.assertEqual(sum(lesson["hours"] for lesson in lessons), quarter_hours)
        self.assertFalse(any(lesson.get("assessment_type") for lesson in lessons))

    def test_docx_penultimate_row_keeps_content_and_has_no_soch(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "soch-empty-content.docx"
            build(data, output)
            document = Document(output)
            quarter = data["quarters"][0]
            lesson = quarter["lessons"][-2]
            row = next(
                row
                for row in document.tables[0].rows
                if row.cells[0].text.strip() == str(lesson["number"])
            )
            self.assertEqual(row.cells[2].text.strip(), lesson["topic"])
            self.assertEqual(row.cells[3].text.strip(), lesson["learning_objectives"])
            self.assertEqual(row.cells[4].text.strip(), "1")
            self.assertFalse(contains_primary_summative_marker(row.cells[6].text.strip()))

    def test_clears_misplaced_and_duplicate_soch_markers(self):
        data = make_data(2)
        lessons = data["quarters"][0]["lessons"]
        lessons[0]["assessment_type"] = "СОЧ"
        lessons[-1]["assessment_type"] = "соч"
        validate(data)
        self.assertFalse(any(lesson.get("assessment_type") for lesson in lessons))

    def test_does_not_insert_soch_when_not_required(self):
        data = make_data(2)
        quarter = data["quarters"][0]
        quarter["soch_required"] = False
        validate(data)
        self.assertFalse(any(
            lesson.get("assessment_type") == "СОЧ"
            for lesson in quarter["lessons"]
        ))

    def test_final_short_section_does_not_require_pre_soch_placement(self):
        data = make_data(1)
        lessons = data["quarters"][0]["lessons"]
        lessons[-2]["section"] = "Қысқа қорытынды бөлім"
        lessons[-1]["section"] = "Қысқа қорытынды бөлім"
        data["cross_cutting_themes"].append("Қысқа қорытынды бөлім")
        validate(data)
        self.assertFalse(any(lesson.get("assessment_type") for lesson in lessons))

    def test_allows_repeated_weekdays_for_multiple_daily_lessons(self):
        validate(make_data(15))

    def test_rejects_zero_hours(self):
        data = make_data(1)
        data["hours_per_week"] = 0
        with self.assertRaisesRegex(ValueError, "greater than or equal to 1"):
            validate(data)

    def test_allows_two_compatible_standard_objectives(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["learning_objectives"] = (
            "2.1.1.1 – описывать признаки объекта; "
            "2.1.1.2 – называть основные части объекта"
        )
        validate(data)

    def test_rejects_three_objectives_in_one_lesson(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["learning_objectives"] = (
            "2.1.1.1 – описывать признаки; 2.1.1.2 – называть части; "
            "2.1.1.3 – объяснять назначение"
        )
        with self.assertRaisesRegex(ValueError, "ең көбі екі мақсат"):
            validate(data)

    def test_rejects_complex_objective_combined_with_another(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["learning_objectives"] = (
            "2.1.1.1 – анализировать причины изменения; "
            "2.1.1.2 – описывать последствия изменения"
        )
        with self.assertRaisesRegex(ValueError, "Күрделі немесе көлемді"):
            validate(data)

    def test_rejects_non_monochrome_theme(self):
        data = make_data(1)
        data["color_theme"] = {"mode": "restrained-color"}
        with self.assertRaisesRegex(ValueError, "Only official monochrome"):
            validate(data)

    def test_preserves_yo_and_kazakh_u_stroke_with_times_new_roman(self):
        data = make_data(1)
        data["subject"] = "Қазақ тілі: Ёлка және Ұлттық мұра"
        data["subject_profile"] = "languages_literature"
        lesson = data["quarters"][0]["lessons"][0]
        lesson["section"] = "Ёж және Ұлы дала"
        data["cross_cutting_themes"].append("Ёж және Ұлы дала")
        lesson["topic"] = "Ёлка мен Ұлттық өрнек"
        lesson["learning_objectives"] = (
            "2.1.1.1 – различать Ё/ё және Ұ/ұ таңбаларын дәл көрсету"
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "unicode.docx"
            build(data, output)
            document = Document(output)
            text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            for char in "ЁёҰұ":
                self.assertIn(char, text)
            font_values = document._element.xpath(".//w:rFonts/@w:ascii")
            self.assertTrue(font_values)
            self.assertEqual(set(font_values), {"Times New Roman"})
            for channel in ("hAnsi", "eastAsia", "cs"):
                values = document._element.xpath(f".//w:rFonts/@w:{channel}")
                self.assertEqual(len(values), len(font_values))
                self.assertEqual(set(values), {"Times New Roman"})

    def test_rejects_unicode_replacement_character(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["topic"] = "Қате � таңба"
        with self.assertRaisesRegex(ValueError, "replacement character U\\+FFFD"):
            validate(data)

    def test_requires_explicit_language_choice_for_ru_kk_mismatch(self):
        data = make_data(1)
        data["requested_language"] = "ru"
        data["objectives_language"] = "kk"
        with self.assertRaisesRegex(ValueError, "КТП қай тілде құрастырылсын"):
            validate(data)

    def test_accepts_explicit_language_choice_after_ru_kk_mismatch(self):
        data = make_data(1)
        data["requested_language"] = "ru"
        data["objectives_language"] = "kk"
        data["language"] = "kk"
        data["language_mismatch_confirmed"] = True
        validate(data)

    def test_holiday_lesson_moves_to_next_lesson_and_is_marked(self):
        data = make_data(1)
        data["non_instruction_dates"] = ["2026-09-08"]
        data["public_holidays"] = [{
            "date": "2026-09-08", "name": "Тестовый праздник", "official_transfer_date": None
        }]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "holiday.docx"
            build(data, output)
            affected = data["quarters"][0]["lessons"][1]
            self.assertEqual(affected["date"], "15.09.2026")
            self.assertIn("Тестовый праздник", affected["note"])
            self.assertIn("08.09.2026", affected["note"])
            self.assertIn("15.09.2026", affected["note"])
            document = Document(output)
            notes = [row.cells[6].text for row in document.tables[0].rows]
            self.assertTrue(any("Тестовый праздник" in value for value in notes))

    def test_holiday_uses_official_transfer_date_first(self):
        data = make_data(1)
        data["non_instruction_dates"] = ["2026-09-08"]
        data["public_holidays"] = [{
            "date": "2026-09-08", "name": "Тестовый праздник",
            "official_transfer_date": "2026-09-12"
        }]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "official-transfer.docx"
            build(data, output)
            affected = data["quarters"][0]["lessons"][1]
            self.assertEqual(affected["date"], "12.09.2026")
            self.assertIn("официальный перенос", affected["note"])

    def test_rejects_holiday_missing_from_non_instruction_dates(self):
        data = make_data(1)
        data["public_holidays"] = [{
            "date": "2026-09-08", "name": "Тестовый праздник", "official_transfer_date": None
        }]
        with self.assertRaisesRegex(ValueError, "must also be in non_instruction_dates"):
            validate(data)

    def test_asks_teacher_when_no_in_quarter_transfer_date_exists(self):
        data = make_data(1)
        data["language"] = "kk"
        data["requested_language"] = "kk"
        data["objectives_language"] = "kk"
        holiday_dates = [
            "2026-09-01", "2026-09-08", "2026-09-15", "2026-09-22",
            "2026-09-29", "2026-10-06", "2026-10-13", "2026-10-20",
        ]
        data["non_instruction_dates"] = holiday_dates
        data["public_holidays"] = [
            {"date": value, "name": "Тест мерекесі", "official_transfer_date": None,
             "teacher_transfer_date": None}
            for value in holiday_dates
        ]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "needs-teacher-date.docx"
            with self.assertRaisesRegex(
                ValueError,
                "Тест мерекесі.*01\\.09\\.2026.*1 четверть.*қай күнге ауыстырайық",
            ):
                build(data, output)
            self.assertFalse(output.exists())

    def test_uses_teacher_confirmed_transfer_date(self):
        data = make_data(1)
        data["non_instruction_dates"] = ["2026-09-08"]
        data["public_holidays"] = [{
            "date": "2026-09-08", "name": "Тест мерекесі",
            "official_transfer_date": None, "teacher_transfer_date": "2026-09-12"
        }]
        data["language"] = "kk"
        data["requested_language"] = "kk"
        data["objectives_language"] = "kk"
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "teacher-transfer.docx"
            build(data, output)
            affected = data["quarters"][0]["lessons"][1]
            self.assertEqual(affected["date"], "12.09.2026")
            self.assertIn("күнді мұғалім растады", affected["note"])

    def test_build_uses_only_white_shading_and_black_text(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "monochrome.docx"
            build(data, output)
            document = Document(output)
            fills = document.tables[0]._tbl.xpath(".//w:shd/@w:fill")
            colors = document._element.xpath(".//w:color/@w:val")
            self.assertTrue(fills)
            self.assertEqual(set(fills), {"FFFFFF"})
            self.assertTrue(colors)
            self.assertEqual(set(colors), {"000000"})

    def test_structure_audit_rejects_colored_shading(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "colored.docx"
            build(data, output)
            document = Document(output)
            cell_properties = document.tables[0].cell(0, 0)._tc.get_or_add_tcPr()
            shading = cell_properties.xpath("./w:shd")[0]
            shading.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", "D9EAF7")
            document.save(output)
            with self.assertRaisesRegex(ValueError, "Colored shading"):
                audit_docx_structure(output, 34, False)

    def test_extended_objectives_is_optional(self):
        data = make_data(2)
        data.pop("extended_objectives")
        validate(data)

    def test_builds_one_continuous_main_table(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "ktp.docx"
            build(data, output)
            document = Document(output)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(len(document.tables[0].rows), 1 + 4 + 34)
            self.assertIsNotNone(document.tables[0]._tbl.tblGrid)
            header_rows = document.tables[0]._tbl.xpath("./w:tr[w:trPr/w:tblHeader]")
            self.assertEqual(len(header_rows), 1)

    def test_build_writes_no_automatic_soch_in_any_quarter(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "ktp.docx"
            build(data, output)
            document = Document(output)
            note_values = [row.cells[6].text.strip() for row in document.tables[0].rows]
            self.assertFalse(any(contains_primary_summative_marker(note) for note in note_values))
            for quarter in data["quarters"]:
                self.assertFalse(any(
                    lesson.get("assessment_type") for lesson in quarter["lessons"]
                ))

    def test_conflicts_add_only_one_appendix_table(self):
        data = make_data(1)
        data["source_conflicts"] = [{
            "location": "I четверть",
            "type": "Неблокирующее расхождение",
            "source_text": "Исходный текст",
            "resolution": "Отмечено в приложении",
        }]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "ktp-with-conflicts.docx"
            build(data, output)
            document = Document(output)
            self.assertEqual(len(document.tables), 2)

    def test_structure_audit_rejects_fragmented_main_table(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "fragmented.docx"
            build(data, output)
            document = Document(output)
            document.add_table(rows=2, cols=7)
            document.save(output)
            with self.assertRaisesRegex(ValueError, "one main KTP table"):
                audit_docx_structure(output, 34, False)

    def test_structure_audit_rejects_manual_page_break(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "manual-break.docx"
            build(data, output)
            document = Document(output)
            document.paragraphs[0].add_run().add_break(WD_BREAK.PAGE)
            document.save(output)
            with self.assertRaisesRegex(ValueError, "Manual page breaks"):
                audit_docx_structure(output, 34, False)

    def test_skill_contract_contains_primary_requirements(self):
        skill_text = (Path(__file__).parents[2] / "references" / "primary-mode.md").read_text(encoding="utf-8")
        required_phrases = [
            "1–4 аралығындағы сынып",
            "education_level` тек `primary",
            "summative_assessment_required",
            "33 апталық",
            "БЖБ/ТЖБ",
            "СОР/СОЧ",
            "soch_required: false",
        ]
        for phrase in required_phrases:
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, skill_text)


if __name__ == "__main__":
    unittest.main()
