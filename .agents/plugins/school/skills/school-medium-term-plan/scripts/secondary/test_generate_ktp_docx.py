#!/usr/bin/env python3
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

from generate_ktp_docx import (
    EXPECTED_QUARTER_DATES,
    EXPECTED_WEEKS,
    audit_docx_structure,
    build,
    summative_assessment_exemption_reason,
    validate,
)


def make_data(hours_per_week: int) -> dict:
    weekdays = [1 + (index % 4) for index in range(hours_per_week)]
    quarters = []
    number = 1
    for index, weeks in enumerate(EXPECTED_WEEKS):
        count = weeks * hours_per_week
        lessons = []
        for position in range(count):
            lessons.append({
                "number": number,
                "section": f"Раздел {index + 1}",
                "topic": "Официальное содержание",
                "learning_objectives": "5.1.1.1 – описывать официальное содержание учебной программы",
                "hours": 1,
                "assessment_type": "",
                "note": "",
            })
            number += 1
        quarters.append({
            "name": f"{index + 1} четверть",
            "weeks": weeks,
            "soch_required": True,
            "start_date": EXPECTED_QUARTER_DATES[index][0],
            "end_date": EXPECTED_QUARTER_DATES[index][1],
            "lessons": lessons,
        })
    return {
        "subject": "Математика",
        "grade": 5,
        "education_level": "basic_secondary",
        "instruction_language": "ru",
        "profile_direction": "not_applicable",
        "curriculum_order": "№399",
        "curriculum_appendix": "52",
        "curriculum_revision_date": "2026-08-14",
        "subject_profile": "mathematics_informatics",
        "language": "ru",
        "requested_language": "ru",
        "objectives_language": "ru",
        "language_mismatch_confirmed": False,
        "academic_year": "2026–2027",
        "total_hours": 34 * hours_per_week,
        "hours_per_week": hours_per_week,
        "lesson_weekdays": weekdays,
        "extended_objectives": [],
        "content_source": "Официальная тестовая программа",
        "content_source_type": "teacher_uploaded_program",
        "teacher_program_file": "teacher-curriculum.pdf",
        "official_program_verified": True,
        "source_content_complete": True,
        "source_conflicts": [],
        "calendar_verified": True,
        "calendar_source_type": "builtin_with_teacher_additions",
        "teacher_calendar_file": "teacher-test-calendar.pdf",
        "teacher_calendar_additions": [{
            "date": "2027-05-27",
            "name": "Тестовое подтверждение переменного праздника",
            "kind": "teacher_confirmed_day_off",
            "resolves_pending": "Құрбан айттың бірінші күні",
            "official_transfer_date": None,
            "teacher_transfer_date": None,
        }],
        "teacher_non_instruction_dates": [],
        "calendar_conflicts": [],
        "non_instruction_dates": [],
        "public_holidays": [],
        "quarters": quarters,
    }


class DynamicHoursValidationTest(unittest.TestCase):
    def test_rejects_external_url_anywhere_in_ktp_input(self):
        data = make_data(1)
        data["calendar_source"] = "https://example.org/calendar"
        with self.assertRaisesRegex(ValueError, "External URLs are forbidden"):
            validate(data)

    def test_rejects_external_url_field_in_nested_calendar_data(self):
        data = make_data(1)
        data["public_holidays"] = [{"url": "external-source"}]
        with self.assertRaisesRegex(ValueError, "External URL fields are forbidden"):
            validate(data)

    def test_rejects_unapproved_calendar_source_type(self):
        data = make_data(1)
        data["calendar_source_type"] = "official_web_search"
        with self.assertRaisesRegex(ValueError, "built-in approved calendar"):
            validate(data)

    def test_pending_mandatory_calendar_data_stops_for_builtin_only(self):
        data = make_data(1)
        data["calendar_source_type"] = "builtin_approved_calendar_2026_2027"
        data["teacher_calendar_file"] = ""
        data["teacher_calendar_additions"] = []
        with self.assertRaisesRegex(ValueError, "request a source from the teacher"):
            validate(data)

    def test_builtin_calendar_replaces_model_supplied_calendar_arrays(self):
        data = make_data(1)
        data["public_holidays"] = []
        data["non_instruction_dates"] = []
        validate(data)
        self.assertIn("2026-12-16", data["non_instruction_dates"])
        self.assertTrue(any(item["date"] == "2026-12-16" for item in data["public_holidays"]))
        self.assertTrue(data["calendar_source_complete"])

    def test_accepts_zero_width_formatting_inside_objective_code(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["learning_objectives"] = (
            "6.\u200b4.\u200b1.\u200b1 описывать свойства музыкального произведения "
            "и сохранять текст\u200bформулировки"
        )

        validate(data)

        normalized = data["quarters"][0]["lessons"][0]["learning_objectives"]
        self.assertTrue(normalized.startswith("6.4.1.1 "))
        self.assertIn("текст\u200bформулировки", normalized)

    def test_summative_exemption_subject_grade_matrix(self):
        cases = [
            ("Музыка", 6),
            ("Көркем еңбек", 9),
            ("Физическая культура", 11),
            ("Светскость и основы религиоведения", 9),
            ("Основы предпринимательства и бизнеса", 10),
            ("Графика и проектирование", 11),
            ("НВП", 10),
        ]
        for subject, grade in cases:
            with self.subTest(subject=subject, grade=grade):
                self.assertIsNotNone(summative_assessment_exemption_reason({
                    "subject": subject,
                    "grade": grade,
                    "academic_year": "2026–2027",
                }))
        self.assertIsNone(summative_assessment_exemption_reason({
            "subject": "Музыка",
            "grade": 7,
            "academic_year": "2026–2027",
        }))

    def test_exempt_subject_clears_sor_and_soch_markers(self):
        data = make_data(1)
        data["subject"] = "Музыка"
        data["subject_profile"] = "arts_technology"
        data["quarters"][0]["lessons"][0]["assessment_type"] = "СОР №9"
        validate(data)
        self.assertTrue(data["summative_assessment_exempt"])
        self.assertTrue(all(not quarter["soch_required"] for quarter in data["quarters"]))
        self.assertFalse(any(
            lesson.get("assessment_type")
            for quarter in data["quarters"]
            for lesson in quarter["lessons"]
        ))
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "music-without-summative-assessment.docx"
            build(data, output)
            document_text = "\n".join(
                cell.text
                for table in Document(output).tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertFalse(any(marker in document_text for marker in ("БЖБ", "ТЖБ", "СОР", "СОЧ")))

    def test_variative_component_clears_summative_markers(self):
        data = make_data(1)
        data["curriculum_component"] = "variative"
        validate(data)
        self.assertTrue(data["summative_assessment_exempt"])
        self.assertFalse(any(
            lesson.get("assessment_type")
            for quarter in data["quarters"]
            for lesson in quarter["lessons"]
        ))

    def test_rejects_exact_objectives_without_teacher_uploaded_program(self):
        data = make_data(1)
        data["content_source_type"] = "teacher_text"
        data["teacher_program_file"] = ""
        with self.assertRaisesRegex(ValueError, "teacher-uploaded curriculum"):
            validate(data)

    def test_uses_only_kazakh_assessment_terms_in_kazakh_ktp(self):
        data = make_data(1)
        data.update({
            "language": "kk",
            "requested_language": "kk",
            "objectives_language": "kk",
            "instruction_language": "kk",
        })
        validate(data)
        lessons = data["quarters"][0]["lessons"]
        self.assertEqual(lessons[-3]["assessment_type"], "БЖБ №1")
        self.assertEqual(lessons[-2]["assessment_type"], "ТЖБ")
        self.assertFalse(any(
            "СОР" in lesson["assessment_type"] or "СОЧ" in lesson["assessment_type"]
            for lesson in lessons
        ))

    def test_uses_only_english_assessment_terms_in_english_ktp(self):
        data = make_data(1)
        data.update({
            "language": "en",
            "requested_language": "en",
            "objectives_language": "en",
            "instruction_language": "en",
        })
        for quarter_index, quarter in enumerate(data["quarters"], start=1):
            quarter["name"] = f"Term {quarter_index}"
            for lesson in quarter["lessons"]:
                lesson["section"] = f"Section {quarter_index}"
                lesson["topic"] = "Official curriculum content"
                lesson["learning_objectives"] = (
                    "5.1.1.1 – describe official curriculum content"
                )
        validate(data)
        lessons = data["quarters"][0]["lessons"]
        self.assertEqual(lessons[-3]["assessment_type"], "SAU №1")
        self.assertEqual(lessons[-2]["assessment_type"], "SAT")
        self.assertFalse(any(
            any(term in lesson["assessment_type"] for term in ("БЖБ", "ТЖБ", "СОР", "СОЧ"))
            for lesson in lessons
        ))

    def test_normalizes_mixed_manual_markers_to_english(self):
        data = make_data(1)
        data.update({
            "language": "en",
            "requested_language": "en",
            "objectives_language": "en",
            "instruction_language": "en",
        })
        data["quarters"][0]["lessons"][0]["assessment_type"] = "СОР №9"
        data["quarters"][0]["lessons"][1]["assessment_type"] = "ТЖБ"
        validate(data)
        lessons = data["quarters"][0]["lessons"]
        self.assertEqual(lessons[-3]["assessment_type"], "SAU №1")
        self.assertEqual(lessons[-2]["assessment_type"], "SAT")

    def test_builds_english_docx_with_english_headers_and_markers(self):
        data = make_data(1)
        data.update({
            "subject": "Mathematics",
            "language": "en",
            "requested_language": "en",
            "objectives_language": "en",
            "instruction_language": "en",
            "content_source": "Official curriculum",
            "calendar_source": "Official academic calendar",
        })
        for quarter_index, quarter in enumerate(data["quarters"], start=1):
            quarter["name"] = f"Term {quarter_index}"
            for lesson in quarter["lessons"]:
                lesson["section"] = f"Section {quarter_index}"
                lesson["topic"] = "Official curriculum content"
                lesson["learning_objectives"] = (
                    "5.1.1.1 – describe official curriculum content"
                )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "english-ktp.docx"
            build(data, output)
            document = Document(output)
            self.assertIn("Medium-term (calendar-thematic) subject plan", document.paragraphs[0].text)
            self.assertEqual(
                [cell.text.strip() for cell in document.tables[0].rows[0].cells],
                ["No.", "Section /\nCross-cutting themes", "Lesson topic", "Learning objectives",
                 "Number of hours", "Dates", "Notes"],
            )
            notes = [row.cells[6].text.strip() for row in document.tables[0].rows]
            self.assertIn("SAU №1", notes)
            self.assertEqual(notes.count("SAT"), 4)

    def test_objective_language_does_not_override_instruction_language(self):
        data = make_data(1)
        data.update({
            "language": "en",
            "requested_language": "en",
            "objectives_language": "kk",
            "instruction_language": "en",
        })
        validate(data)

    def test_accepts_english_after_explicit_mismatch_choice(self):
        data = make_data(1)
        data.update({
            "language": "en",
            "requested_language": "en",
            "objectives_language": "en",
            "instruction_language": "en",
            "language_mismatch_confirmed": True,
        })
        validate(data)
        self.assertEqual(data["quarters"][0]["lessons"][-2]["assessment_type"], "SAT")

    def test_rejects_subject_profile_mismatch(self):
        data = make_data(1)
        data["subject_profile"] = "natural_sciences"
        with self.assertRaisesRegex(ValueError, "subject_profile must be"):
            validate(data)

    def test_requires_profile_direction_for_grade_10_subject(self):
        data = make_data(1)
        data["grade"] = 10
        data["education_level"] = "general_secondary"
        data["subject"] = "Физика"
        data["subject_profile"] = "natural_sciences"
        with self.assertRaisesRegex(ValueError, "ЖМБ немесе ҚГБ"):
            validate(data)

    def test_accepts_profile_direction_for_grade_10_subject(self):
        data = make_data(1)
        data["grade"] = 10
        data["education_level"] = "general_secondary"
        data["subject"] = "Физика"
        data["subject_profile"] = "natural_sciences"
        data["profile_direction"] = "ЖМБ"
        validate(data)

    def test_requires_safety_metadata_for_laboratory_lesson(self):
        data = make_data(1)
        data["subject"] = "Физика"
        data["subject_profile"] = "natural_sciences"
        lesson = data["quarters"][0]["lessons"][0]
        lesson["lesson_mode"] = "laboratory"
        lesson["practical_work"] = "Тығыздықты анықтау"
        with self.assertRaisesRegex(ValueError, "safety_required"):
            validate(data)

    def test_accepts_complete_laboratory_metadata(self):
        data = make_data(1)
        data["subject"] = "Физика"
        data["subject_profile"] = "natural_sciences"
        lesson = data["quarters"][0]["lessons"][0]
        lesson.update({
            "lesson_mode": "laboratory",
            "practical_work": "Тығыздықты анықтау",
            "safety_required": True,
        })
        validate(data)

    def test_accepts_dynamic_hours_without_upper_limit(self):
        for hours_per_week in (1, 2, 10, 11, 15):
            with self.subTest(hours_per_week=hours_per_week):
                validate(make_data(hours_per_week))

    def test_rejects_fixed_68_hours_for_one_hour_subject(self):
        data = make_data(1)
        data["total_hours"] = 68
        with self.assertRaisesRegex(ValueError, "do not equal total hours"):
            validate(data)

    def test_normalizes_manual_sor_to_section_end_label(self):
        data = make_data(2)
        data["quarters"][0]["lessons"][0]["assessment_type"] = "СОР"
        validate(data)
        lessons = data["quarters"][0]["lessons"]
        self.assertEqual(lessons[0]["assessment_type"], "")
        self.assertEqual(lessons[-3]["assessment_type"], "СОР №1")
        self.assertEqual(lessons[-2]["assessment_type"], "СОЧ")
        self.assertEqual(lessons[-1]["assessment_type"], "")

    def test_numbers_section_assessments_within_each_quarter(self):
        data = make_data(1)
        lessons = data["quarters"][0]["lessons"]
        for index, lesson in enumerate(lessons):
            lesson["section"] = "Бөлім A" if index < 3 else "Бөлім B"
        validate(data)
        self.assertEqual(lessons[2]["assessment_type"], "СОР №1")
        self.assertEqual(lessons[-3]["assessment_type"], "СОР №2")

    def test_docx_displays_exact_section_assessment_labels(self):
        data = make_data(1)
        lessons = data["quarters"][0]["lessons"]
        for index, lesson in enumerate(lessons):
            lesson["section"] = "Бөлім A" if index < 3 else "Бөлім B"
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "section-assessments.docx"
            build(data, output)
            document = Document(output)
            notes = [row.cells[6].text.strip() for row in document.tables[0].rows]
            self.assertIn("СОР №1", notes)
            self.assertIn("СОР №2", notes)

    def test_section_assessment_numbering_restarts_each_quarter(self):
        data = make_data(1)
        validate(data)
        for quarter in data["quarters"]:
            self.assertEqual(
                quarter["lessons"][-3]["assessment_type"],
                "СОР №1",
            )

    def test_section_continuing_across_quarter_is_not_marked_at_boundary(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][-1]["section"] = "Жалғасатын бөлім"
        data["quarters"][1]["lessons"][0]["section"] = "Жалғасатын бөлім"
        validate(data)
        self.assertNotIn("СОР", data["quarters"][0]["lessons"][-1]["assessment_type"])
        self.assertEqual(
            data["quarters"][1]["lessons"][0]["assessment_type"],
            "СОР №1",
        )

    def test_moves_section_assessment_before_soch(self):
        data = make_data(1)
        lessons = data["quarters"][0]["lessons"]
        lessons[-1]["section"] = "Келесі бөлім"
        data["quarters"][1]["lessons"][0]["section"] = "Келесі бөлім"
        validate(data)
        self.assertEqual(
            lessons[-3]["assessment_type"],
            "СОР №1",
        )
        self.assertEqual(lessons[-2]["assessment_type"], "СОЧ")
        self.assertEqual(lessons[-1]["assessment_type"], "")

    def test_automatically_inserts_soch_on_penultimate_lesson(self):
        data = make_data(2)
        lessons = data["quarters"][0]["lessons"]
        lesson_count = len(lessons)
        quarter_hours = sum(lesson["hours"] for lesson in lessons)
        validate(data)
        self.assertEqual(lessons[-2]["assessment_type"], "СОЧ")
        self.assertEqual(lessons[-2]["topic"], "")
        self.assertEqual(lessons[-2]["learning_objectives"], "")
        self.assertEqual(lessons[-2]["hours"], 1)
        self.assertEqual(len(lessons), lesson_count)
        self.assertEqual(sum(lesson["hours"] for lesson in lessons), quarter_hours)
        self.assertFalse(any(
            lesson.get("assessment_type") == "СОЧ"
            for lesson in lessons[:-2] + lessons[-1:]
        ))

    def test_docx_soch_row_contains_only_assessment_not_topic_or_objective(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "soch-empty-content.docx"
            build(data, output)
            document = Document(output)
            quarter = data["quarters"][0]
            lesson = quarter["lessons"][-2]
            row = next(
                row
                for row in document.tables[0].rows
                if row.cells[0].text.strip() == str(lesson["number"])
            )
            self.assertEqual(row.cells[2].text.strip(), "")
            self.assertEqual(row.cells[3].text.strip(), "")
            self.assertEqual(row.cells[4].text.strip(), "1")
            self.assertEqual(row.cells[6].text.strip(), "СОЧ")

    def test_moves_misplaced_and_duplicate_soch_to_penultimate_lesson(self):
        data = make_data(2)
        lessons = data["quarters"][0]["lessons"]
        lessons[0]["assessment_type"] = "СОЧ"
        lessons[-1]["assessment_type"] = "соч"
        validate(data)
        positions = [
            index for index, lesson in enumerate(lessons)
            if lesson.get("assessment_type") == "СОЧ"
        ]
        self.assertEqual(positions, [len(lessons) - 2])

    def test_does_not_insert_soch_when_not_required(self):
        data = make_data(2)
        quarter = data["quarters"][0]
        quarter["soch_required"] = False
        validate(data)
        self.assertFalse(any(
            lesson.get("assessment_type") == "СОЧ"
            for lesson in quarter["lessons"]
        ))

    def test_stops_when_final_section_has_no_pre_soch_lesson(self):
        data = make_data(1)
        lessons = data["quarters"][0]["lessons"]
        lessons[-2]["section"] = "Қысқа қорытынды бөлім"
        lessons[-1]["section"] = "Қысқа қорытынды бөлім"
        with self.assertRaisesRegex(ValueError, "қай сабаққа орналастырайық"):
            validate(data)

    def test_allows_repeated_weekdays_for_multiple_daily_lessons(self):
        validate(make_data(15))

    def test_rejects_zero_hours(self):
        data = make_data(1)
        data["hours_per_week"] = 0
        with self.assertRaisesRegex(ValueError, "greater than or equal to 1"):
            validate(data)

    def test_allows_two_compatible_standard_objectives(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["learning_objectives"] = (
            "5.1.1.1 – описывать признаки объекта; "
            "5.1.1.2 – называть основные части объекта"
        )
        validate(data)

    def test_rejects_three_objectives_in_one_lesson(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["learning_objectives"] = (
            "5.1.1.1 – описывать признаки; 5.1.1.2 – называть части; "
            "5.1.1.3 – объяснять назначение"
        )
        with self.assertRaisesRegex(ValueError, "ең көбі екі мақсат"):
            validate(data)

    def test_rejects_complex_objective_combined_with_another(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["learning_objectives"] = (
            "5.1.1.1 – анализировать причины изменения; "
            "5.1.1.2 – описывать последствия изменения"
        )
        with self.assertRaisesRegex(ValueError, "Күрделі немесе көлемді"):
            validate(data)

    def test_rejects_non_monochrome_theme(self):
        data = make_data(1)
        data["color_theme"] = {"mode": "restrained-color"}
        with self.assertRaisesRegex(ValueError, "Only official monochrome"):
            validate(data)

    def test_preserves_yo_and_kazakh_u_stroke_with_times_new_roman(self):
        data = make_data(1)
        data["subject"] = "Қазақ тілі: Ёлка және Ұлттық мұра"
        data["subject_profile"] = "languages_literature"
        lesson = data["quarters"][0]["lessons"][0]
        lesson["section"] = "Ёж және Ұлы дала"
        lesson["topic"] = "Ёлка мен Ұлттық өрнек"
        lesson["learning_objectives"] = (
            "5.1.1.1 – различать Ё/ё және Ұ/ұ таңбаларын дәл көрсету"
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "unicode.docx"
            build(data, output)
            document = Document(output)
            text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            for char in "ЁёҰұ":
                self.assertIn(char, text)
            font_values = document._element.xpath(".//w:rFonts/@w:ascii")
            self.assertTrue(font_values)
            self.assertEqual(set(font_values), {"Times New Roman"})
            for channel in ("hAnsi", "eastAsia", "cs"):
                values = document._element.xpath(f".//w:rFonts/@w:{channel}")
                self.assertEqual(len(values), len(font_values))
                self.assertEqual(set(values), {"Times New Roman"})

    def test_rejects_unicode_replacement_character(self):
        data = make_data(1)
        data["quarters"][0]["lessons"][0]["topic"] = "Қате � таңба"
        with self.assertRaisesRegex(ValueError, "replacement character U\\+FFFD"):
            validate(data)

    def test_objective_language_mismatch_does_not_trigger_document_language_question(self):
        data = make_data(1)
        data["requested_language"] = "ru"
        data["objectives_language"] = "kk"
        validate(data)

    def test_accepts_explicit_document_language_override(self):
        data = make_data(1)
        data["requested_language"] = "ru"
        data["objectives_language"] = "kk"
        data["language"] = "kk"
        data["language_mismatch_confirmed"] = True
        validate(data)

    def test_rejects_unconfirmed_document_language_override(self):
        data = make_data(1)
        data["language"] = "kk"
        with self.assertRaisesRegex(ValueError, "Output language must match"):
            validate(data)

    def test_holiday_lesson_moves_to_next_lesson_and_is_marked(self):
        data = make_data(1)
        data["teacher_calendar_additions"].append({
            "date": "2026-09-08", "name": "Тестовый праздник", "official_transfer_date": None
        })
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "holiday.docx"
            build(data, output)
            affected = data["quarters"][0]["lessons"][1]
            self.assertEqual(affected["date"], "15.09.2026")
            self.assertIn("Тестовый праздник", affected["note"])
            self.assertIn("08.09.2026", affected["note"])
            self.assertIn("15.09.2026", affected["note"])
            document = Document(output)
            notes = [row.cells[6].text for row in document.tables[0].rows]
            self.assertTrue(any("Тестовый праздник" in value for value in notes))

    def test_holiday_uses_official_transfer_date_first(self):
        data = make_data(1)
        data["teacher_calendar_additions"].append({
            "date": "2026-09-08", "name": "Тестовый праздник",
            "official_transfer_date": "2026-09-12"
        })
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "official-transfer.docx"
            build(data, output)
            affected = data["quarters"][0]["lessons"][1]
            self.assertEqual(affected["date"], "12.09.2026")
            self.assertIn("официальный перенос", affected["note"])

    def test_rejects_teacher_addition_replacing_builtin_date(self):
        data = make_data(1)
        data["teacher_calendar_additions"].append({
            "date": "2026-12-16", "name": "Подмена встроенного праздника"
        })
        with self.assertRaisesRegex(ValueError, "must not replace a built-in date"):
            validate(data)

    def test_asks_teacher_when_no_in_quarter_transfer_date_exists(self):
        data = make_data(1)
        data["language"] = "kk"
        data["requested_language"] = "kk"
        data["objectives_language"] = "kk"
        data["instruction_language"] = "kk"
        holiday_dates = [
            "2026-09-01", "2026-09-08", "2026-09-15", "2026-09-22",
            "2026-09-29", "2026-10-06", "2026-10-13", "2026-10-20",
        ]
        data["teacher_calendar_additions"].extend(
            {"date": value, "name": "Тест мерекесі", "official_transfer_date": None,
             "teacher_transfer_date": None}
            for value in holiday_dates
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "needs-teacher-date.docx"
            with self.assertRaisesRegex(
                ValueError,
                "Тест мерекесі.*01\\.09\\.2026.*1 четверть.*қай күнге ауыстырайық",
            ):
                build(data, output)
            self.assertFalse(output.exists())

    def test_uses_teacher_confirmed_transfer_date(self):
        data = make_data(1)
        data["teacher_calendar_additions"].append({
            "date": "2026-09-08", "name": "Тест мерекесі",
            "official_transfer_date": None, "teacher_transfer_date": "2026-09-12"
        })
        data["language"] = "kk"
        data["requested_language"] = "kk"
        data["objectives_language"] = "kk"
        data["instruction_language"] = "kk"
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "teacher-transfer.docx"
            build(data, output)
            affected = data["quarters"][0]["lessons"][1]
            self.assertEqual(affected["date"], "12.09.2026")
            self.assertIn("күнді мұғалім растады", affected["note"])

    def test_build_uses_only_white_shading_and_black_text(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "monochrome.docx"
            build(data, output)
            document = Document(output)
            fills = document.tables[0]._tbl.xpath(".//w:shd/@w:fill")
            colors = document._element.xpath(".//w:color/@w:val")
            self.assertTrue(fills)
            self.assertEqual(set(fills), {"FFFFFF"})
            self.assertTrue(colors)
            self.assertEqual(set(colors), {"000000"})

    def test_structure_audit_rejects_colored_shading(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "colored.docx"
            build(data, output)
            document = Document(output)
            cell_properties = document.tables[0].cell(0, 0)._tc.get_or_add_tcPr()
            shading = cell_properties.xpath("./w:shd")[0]
            shading.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", "D9EAF7")
            document.save(output)
            with self.assertRaisesRegex(ValueError, "Colored shading"):
                audit_docx_structure(output, 34, False)

    def test_extended_objectives_is_optional(self):
        data = make_data(2)
        data.pop("extended_objectives")
        validate(data)

    def test_builds_one_continuous_main_table(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "ktp.docx"
            build(data, output)
            document = Document(output)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(len(document.tables[0].rows), 1 + 4 + 34)
            self.assertIsNotNone(document.tables[0]._tbl.tblGrid)
            header_rows = document.tables[0]._tbl.xpath("./w:tr[w:trPr/w:tblHeader]")
            self.assertEqual(len(header_rows), 1)

    def test_build_writes_one_automatic_soch_per_quarter(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "ktp.docx"
            build(data, output)
            document = Document(output)
            note_values = [row.cells[6].text.strip() for row in document.tables[0].rows]
            self.assertEqual(note_values.count("СОЧ"), 4)
            for quarter in data["quarters"]:
                positions = [
                    index for index, lesson in enumerate(quarter["lessons"])
                    if lesson.get("assessment_type") == "СОЧ"
                ]
                self.assertEqual(positions, [len(quarter["lessons"]) - 2])

    def test_conflicts_add_only_one_appendix_table(self):
        data = make_data(1)
        data["source_conflicts"] = [{
            "location": "I четверть",
            "type": "Неблокирующее расхождение",
            "source_text": "Исходный текст",
            "resolution": "Отмечено в приложении",
        }]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "ktp-with-conflicts.docx"
            build(data, output)
            document = Document(output)
            self.assertEqual(len(document.tables), 2)

    def test_structure_audit_rejects_fragmented_main_table(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "fragmented.docx"
            build(data, output)
            document = Document(output)
            document.add_table(rows=2, cols=7)
            document.save(output)
            with self.assertRaisesRegex(ValueError, "one main KTP table"):
                audit_docx_structure(output, 34, False)

    def test_structure_audit_rejects_manual_page_break(self):
        data = make_data(1)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "manual-break.docx"
            build(data, output)
            document = Document(output)
            document.paragraphs[0].add_run().add_break(WD_BREAK.PAGE)
            document.save(output)
            with self.assertRaisesRegex(ValueError, "Manual page breaks"):
                audit_docx_structure(output, 34, False)

    def test_skill_contract_contains_onboarding_gate(self):
        skill_text = (Path(__file__).parents[2] / "references" / "secondary-mode.md").read_text(encoding="utf-8")
        required_phrases = [
            "I just added",
            "make up a realistic user prompt",
            "остановиться и дождаться реального ответа",
            "не создавать КТП/DOCX в этом же ходе",
            "не более двух пользовательских циклов",
            "Не задавать отдельный вопрос о языке документа",
            "СОР №1",
            "СОР №2",
        ]
        for phrase in required_phrases:
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, skill_text)


if __name__ == "__main__":
    unittest.main()
